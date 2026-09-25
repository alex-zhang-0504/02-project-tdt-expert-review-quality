"""One vote per manager; editable Excel task exchange and atomic merge."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from .assessment import tasks, manager_task, is_excluded
from .questionnaire import dimensions, snapshot as questionnaire_snapshot
from .subjective import UNJUDGED, RULE_VERSION, ReviewInput, save_review


def task_scope(manager):
    return sha256(json.dumps(manager['experts'], sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def average_review(analysis, expert_name, policy):
    from .score_statistics import questionnaire_score
    catalog = tasks(analysis)
    expected = [m for m in catalog['managers'] if expert_name in m['experts'] and not is_excluded(analysis, m['manager_id'], expert_name)]
    answers = [questionnaire_score(analysis.manager_reviews.get(m['manager_id'], {}).get(expert_name, {}), policy, questionnaire_snapshot(analysis))[0] for m in expected]
    completed = sum(all(i['responded'] for i in row) for row in answers)
    items = []
    for index, dimension in enumerate(dimensions(analysis)):
        values = [row[index]['score'] for row in answers if row[index]['score'] is not None]
        contribution = dimension['id'] == 'contribution'
        items.append({'dimension': dimension['title'], 'option': '有效最高分' if contribution else '逐题有效等权平均',
                      'evidence_missing': any(row[index]['evidence_missing'] for row in answers),
                      'valid_count': len(values), 'responded_count': sum(row[index]['responded'] for row in answers),
                      'score': (max(values) if contribution else sum(values)/len(values)) if values else None})
    mean = sum(i['score'] for i in items) if all(i['score'] is not None for i in items) else None
    final = bool(expected) and completed == len(expected) and mean is not None and not catalog['unresolved_reports']
    return items, mean if final else None, {'expected': len(expected), 'completed': completed, 'provisional': mean, 'final': final}


def export_task(analysis, manager_id):
    catalog = tasks(analysis)
    manager = next((m for m in catalog['managers'] if m['manager_id'] == manager_id), None)
    if not manager:
        raise ValueError('没有此项目经理的评价任务')
    if catalog['unresolved_reports']:
        raise ValueError('仍有项目经理待识别，不能分发任务')
    wb = Workbook()
    sheet = wb.active
    sheet.title = '经理问卷'
    sheet.append(['评审人', '维度编号', '维度', '档位', '项目编码', '事实依据（100字）', '无法判断或无机会原因', '题干', '职责边界', '行为解释'])
    extra = wb.create_sheet('补充依据')
    extra.append(['评审人', '维度编号', '项目编码', '事实依据（每条100字）'])
    bases = {}
    for name, projects in manager['experts'].items():
        if is_excluded(analysis, manager_id, name):
            continue
        review = analysis.manager_reviews.get(manager_id, {}).get(name, {})
        bases[name] = review.get('revision', 0)
        for dimension in dimensions(analysis):
            rating = review.get('ratings', {}).get(dimension['id'], {})
            choices = dimension['options'] + ([{'id': 'unable', 'title': dimension['unable_title']}] if dimension['id'] != 'contribution' else [])
            label = next((o['title'] for o in choices if o['id'] == rating.get('option')), '')
            sheet.append([name, dimension['id'], dimension['title'], label, rating.get('project_code', ''), rating.get('note', ''), rating.get('reason', ''), dimension['prompt'], dimension['boundary'], '；'.join(o['title']+'：'+o.get('description','') for o in choices)])
            for item in rating.get('evidence', []):
                if item.get('project_code') or item.get('note'): extra.append([name, dimension['id'], item.get('project_code', ''), item.get('note', '')])
            validation = DataValidation(type='list', formula1='"' + ','.join(o['title'] for o in choices) + '"', allow_blank=True)
            validation.errorTitle, validation.error, validation.showErrorMessage = '档位无效', '请使用下拉列表中的档位', True
            sheet.add_data_validation(validation)
            validation.add(sheet.cell(sheet.max_row, 4))
    info = wb.create_sheet('填写说明')
    for row in [['项目经理', manager['name']], ['经理身份', manager_id], ['批次', analysis.assessment['batch_id']],
                ['操作', '选择中文行为选项，填写项目、事实或无法判断原因；更多证据填写补充依据。空白表示待评价。'],
                ['回收', '前五题逐题有效等权平均，贡献取有效最高分；无判断状态不入分母。修改同一任务后可重新导入，不增加票数；旧任务不能覆盖更新后的评价。'],
                ['配置指纹', analysis.assessment['policy']['sha256']]]:
        info.append(row)
    for name, codes in manager['experts'].items():
        info.append([name + '可关联项目', '；'.join(codes)])
    for d in dimensions(analysis):
        for o in d['options']:
            info.append([d['title'] + '/' + o['title'], o['title'] + '：' + o['description']])
    meta = wb.create_sheet('_task')
    meta.append([json.dumps({'kind': 'manager-questionnaire-v0.1', 'batch': analysis.assessment['batch_id'],
        'roster_hash': analysis.assessment['roster_hash'], 'policy_hash': analysis.assessment['policy']['sha256'],
        'questionnaire_version': RULE_VERSION, 'questionnaire_hash': questionnaire_snapshot(analysis)['sha256'], 'manager_id': manager_id, 'bases': bases, 'task_scope': task_scope(manager)}, ensure_ascii=False)])
    meta.sheet_state = 'hidden'
    revision = wb.create_sheet('提交修订')
    revision.append(['修订号', 1])
    revision.append(['说明', '每次修改后重新提交，将B1增加1；重复导入同一修订不会重复计票。'])
    evidence = wb.create_sheet('共同项目事实')
    evidence.append(['评审人', '项目', '阶段', '参评', '会签', '意见', '对策判定'])
    for expert in analysis.experts:
        if expert.expert_name not in bases: continue
        for session in expert.sessions:
            if session.project_code not in manager['experts'][expert.expert_name]: continue
            for opinion in session.opinions or [None]:
                evidence.append([expert.expert_name, session.project_code, session.stage, session.attendance,
                    session.signoff, opinion.text if opinion else '无', opinion.ai_status if opinion else ''])
    for row in evidence:
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = 's'
    for ws in [sheet, info, extra]:
        ws.freeze_panes = 'A2'
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 28 if col[0].column < 6 else 60
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def import_task(analysis, content):
    try:
        wb = load_workbook(BytesIO(content), data_only=False)
        meta = json.loads(wb['_task']['A1'].value)
        if not (meta['kind'] == 'manager-questionnaire-v0.1'): raise ValueError()
        if meta.get('questionnaire_version') != RULE_VERSION: raise ValueError()
        if meta.get('questionnaire_hash') != questionnaire_snapshot(analysis)['sha256']: raise ValueError()
        if not (meta['batch'] == analysis.assessment['batch_id']): raise ValueError()
        if not (meta['roster_hash'] == analysis.assessment['roster_hash']): raise ValueError()
        if not (meta['policy_hash'] == analysis.assessment['policy']['sha256']): raise ValueError()
    except (KeyError, TypeError, ValueError, AssertionError):
        raise ValueError('任务文件或批次、名单、评分参数不一致，未导入') from None
    mid = meta['manager_id']
    revision = wb['提交修订']['B1'].value
    if type(revision) is not int or revision < 1:
        raise ValueError('提交修订号必须为正整数')
    trial = deepcopy(analysis)
    grouped, seen = {}, set()
    for row in wb['经理问卷'].iter_rows(min_row=2):
        if any(c.data_type == 'f' for c in row): raise ValueError('问卷不支持公式，请填写档位和文字')
        name, dim, _, option, project, note = [str(c.value or '').strip() for c in row[:6]]
        if (name, dim) in seen: raise ValueError('问卷出现重复维度，未导入')
        seen.add((name, dim))
        if name not in meta['bases'] or dim not in {d['id'] for d in dimensions(analysis)}: raise ValueError('问卷任务或维度被改动，未导入')
        ratings = grouped.setdefault(name, {})
        definition = next(d for d in dimensions(analysis) if d['id'] == dim)
        option = next((o['id'] for o in definition['options'] if o['title'] == option), 'unable' if option==definition['unable_title'] else next((k for k, v in UNJUDGED.items() if v == option), option))
        if option:
            ratings[dim] = dict(option=option, project_code=project, note=note)
            if len(row)>6 and row[6].value: ratings[dim]['reason'] = str(row[6].value).strip()
    if set(grouped) != set(meta['bases']) or len(seen) != len(grouped) * len(dimensions(analysis)):
        raise ValueError('问卷缺少任务或维度行，未导入')
    if '补充依据' in wb:
        for row in wb['补充依据'].iter_rows(min_row=2):
            if any(c.data_type == 'f' for c in row): raise ValueError('依据不支持公式')
            name, dim, project, note = [str(c.value or '').strip() for c in row[:4]]
            if not any((name, dim, project, note)): continue
            if name not in grouped or dim not in grouped[name]: raise ValueError('补充依据须对应已选择的问卷题目')
            grouped[name][dim].setdefault('evidence', []).append({'project_code': project, 'note': note})
    updated = 0
    for name, ratings in grouped.items():
        manager = manager_task(analysis, mid, name)
        if meta.get('task_scope') != task_scope(manager):
            raise ValueError('经理与项目的对应关系已变化，请重新导出任务')
        current = analysis.manager_reviews.get(mid, {}).get(name, {})
        if current.get('ratings') == ratings:
            continue
        base = meta['bases'][name]
        if current and current.get('revision', 0) != base and current.get('task_base') != base:
            raise ValueError('存在更新的评价，请重新导出该经理任务后填写')
        if current.get('task_base') == base and revision <= current.get('task_revision', 0):
            raise ValueError('该任务修订号不高于已接收版本，请增加修订号后提交')
        record = save_review(trial, ReviewInput(analysis_id=analysis.analysis_id, expert_name=name,
            manager_id=mid, evaluator=manager['name'], ratings=ratings))
        record['task_base'] = base
        record['task_revision'] = revision
        updated += 1
    analysis.manager_reviews = trial.manager_reviews
    return {'updated': updated, 'message': f'已更新{updated}份问卷，同一经理仍为一票'}
