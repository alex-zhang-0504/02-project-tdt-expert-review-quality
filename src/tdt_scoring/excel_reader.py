from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900, from_excel
from openpyxl.worksheet.worksheet import Worksheet

from .models import OpinionSource, ProblemRecord, ReviewSession, SignoffRecord, ValidationIssue
from .validation import (
    PROBLEM_NUMBER_PATTERN,
    VALID_CONCLUSIONS,
    validate_problem_number_conflicts,
    validate_session,
    validate_stage_conflicts,
)


SECTION_BASIC = "评审会基本信息列表"
SECTION_SIGNOFF = "评审结论和会签列表"
SECTION_PROBLEMS = "评审问题&建议汇总表"
SECTION_PROBLEMS_V04 = "评审问题汇总表"
SECTION_TITLES = (SECTION_BASIC, SECTION_SIGNOFF, SECTION_PROBLEMS)
SIGNOFF_HEADERS = ("评审角色", "评审人", "参会状况", "会签结果", "评审依据")
PROBLEM_HEADERS = (
    "序号",
    "评审人",
    "问题描述",
    "建议改善措施",
    "验证方式与通过条件",
    "进展",
    "问题状态",
)

PROJECT_CODE_PATTERN = re.compile(r"[（(]([^（）()]*)[）)]")
PROXY_PATTERN = re.compile(r"^(.*?)[（(]\s*(?:代理\s*[:：]?\s*)?(.*?)[）)]\s*$")
ABSENT_ROLE_PATTERN = re.compile(r"^(.*?)[（(][^（）()]+[）)]$")
RECORDED_NO_CONCLUSION_MARKERS = {"-", "－", "—", "–"}
NUMBERED_ITEM_PATTERN = re.compile(
    r"(?m)^[ \t]*(?:\d+[、．.]|[（(]\d+[）)]|[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳])[ \t]*"
)
V04_SECTION_ALIASES = {
    "signoff": (SECTION_SIGNOFF,),
    "problems": (SECTION_PROBLEMS_V04,),
}
V04_BASIC_FIELD_ALIASES = {
    "project_identity": ("技术项目名和编码",),
    "absent_reviewers": ("缺席评审人姓名", "缺席人姓名和角色"),
    "meeting_date": ("TDR会议日期",),
    "stage": ("评审阶段",),
    "meeting_conclusion": ("评审结论",),
}
V04_BASIC_FIELD_NAMES = {
    "project_identity": "技术项目名和编码",
    "absent_reviewers": "缺席评审人姓名",
    "meeting_date": "TDR会议日期",
    "stage": "评审阶段",
    "meeting_conclusion": "评审结论",
}
V04_SIGNOFF_HEADER_ALIASES = {
    "role": ("评审角色",),
    "reviewer": ("评审人姓名", "评审人"),
    "conclusion": ("会签结果",),
    "opinion": ("评审意见",),
}
V04_SIGNOFF_HEADER_NAMES = {
    "role": "评审角色",
    "reviewer": "评审人姓名",
    "conclusion": "会签结果",
    "opinion": "评审意见",
}
V04_PROBLEM_HEADER_ALIASES = {
    "number": ("序号",),
    "reviewers": ("评审人",),
    "description": ("问题描述",),
    "action": ("是否要修正",),
    "verification": ("问题等级",),
    "progress": ("反馈/修改说明", "反馈／修改说明"),
    "status": ("问题状态",),
}
V04_PROBLEM_HEADER_NAMES = {
    "number": "序号",
    "reviewers": "评审人",
    "description": "问题描述",
    "action": "是否要修正",
    "verification": "问题等级",
    "progress": "反馈/修改说明",
    "status": "问题状态",
}
V04_PROJECT_CODE_PATTERN = re.compile(r"^(?=.*\d)[A-Za-z0-9]+(?:[._/][A-Za-z0-9]+)*$")


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ")
    return " ".join(text.split()).strip()


def normalize_label(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"\s+", "", text).strip()


def normalize_parentheses(value: object) -> str:
    return normalize_text(value).translate(str.maketrans({"(": "（", ")": "）", ":": "："}))


def canonical_attendance(value: object) -> str:
    return normalize_parentheses(value)


def canonical_conclusion(value: object) -> tuple[str, bool]:
    raw = normalize_parentheses(value)
    overdue = raw.endswith("（逾期）")
    base = raw[: -len("（逾期）")].strip() if overdue else raw
    aliases = {
        "go": "Go",
        "go with risk": "Go with Risk",
        "redirect": "Redirect",
        "tbd": "TBD",
    }
    return aliases.get(base.casefold(), base), overdue


def canonical_meeting_conclusion(value: object) -> str:
    raw = normalize_parentheses(value)
    aliases = {
        "go": "Go",
        "go with risk": "Go with Risk",
        "redirect": "Redirect",
        "tbd": "TBD",
    }
    return aliases.get(raw.casefold(), raw)


def canonical_problem_status(value: object) -> str:
    raw = normalize_text(value)
    aliases = {
        "open": "open",
        "开启": "open",
        "未关闭": "open",
        "closed": "closed",
        "close": "closed",
        "已关闭": "closed",
    }
    return aliases.get(raw.casefold(), raw)


def parse_reviewer(value: object) -> tuple[str, str | None, str]:
    raw = normalize_text(value)
    match = PROXY_PATTERN.match(raw)
    if not match:
        return raw, None, raw
    expert_name = normalize_text(match.group(1))
    proxy_name = normalize_text(match.group(2))
    return expert_name, proxy_name or None, raw


def split_reviewers(value: object) -> list[str]:
    text = str(value or "")
    if not text:
        return []
    return [
        normalize_text(reviewer)
        for reviewer in re.split(r"[、，,；;／/\r\n]", text)
        if normalize_text(reviewer)
    ]


def split_absent_reviewers(value: object) -> list[str]:
    text = normalize_text(value)
    if not text or text.casefold() in {"无", "-", "none"}:
        return []
    reviewers: list[str] = []
    for item in re.split(r"[、，,；;／/\n]", str(value)):
        normalized = normalize_text(item)
        if not normalized:
            continue
        match = ABSENT_ROLE_PATTERN.match(normalized)
        reviewers.append(normalize_text(match.group(1)) if match else normalized)
    return reviewers


def split_numbered_items(value: object) -> list[str]:
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    matches = list(NUMBERED_ITEM_PATTERN.finditer(text))
    if len(matches) < 2:
        return []
    items: list[str] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        item = normalize_text(text[match.end() : end])
        if item:
            items.append(item)
    return items if len(items) >= 2 else []


def split_opinion_items(value: object) -> list[str]:
    numbered = split_numbered_items(value)
    if numbered:
        return numbered
    lines = [normalize_text(line) for line in str(value or "").splitlines()]
    lines = [line for line in lines if line]
    if len(lines) > 1:
        return lines
    text = normalize_text(value)
    return [text] if text and text != "-" else []


def normalize_opinion_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", normalize_text(value)).casefold()
    text = re.sub(r"^(?:\d+[、.]|\(\d+\)|[①②③④⑤⑥⑦⑧⑨⑩])\s*", "", text)
    return re.sub(r"[\s,，。.;；:：、!?！？()（）\[\]【】\-—–－]", "", text)


def _add_opinion_source(
    sources: list[OpinionSource],
    text: str,
    cell_reference: str,
    raw_text: str,
) -> None:
    key = normalize_opinion_key(text)
    if not key:
        return
    for source in sources:
        if normalize_opinion_key(source.text) == key:
            if cell_reference not in source.cell_references:
                source.cell_references.append(cell_reference)
            if raw_text not in source.raw_texts:
                source.raw_texts.append(raw_text)
            return
    sources.append(
        OpinionSource(
            text=text,
            cell_references=[cell_reference],
            raw_texts=[raw_text],
        )
    )


def parse_project(value: object) -> tuple[str, str]:
    text = normalize_text(value)
    matches = list(PROJECT_CODE_PATTERN.finditer(text))
    if not matches:
        return text, ""
    match = matches[-1]
    project_code = normalize_text(match.group(1))
    project_name = normalize_text(text[: match.start()] + text[match.end() :])
    return project_name, project_code


def parse_v04_project(value: object) -> tuple[str, str]:
    text = normalize_text(value)
    matches = list(re.finditer(r"[-－—–]", text))
    if not matches:
        return text, ""
    boundary = matches[-1]
    project_name = normalize_text(text[: boundary.start()])
    project_code = normalize_text(text[boundary.end() :])
    if not project_name or not V04_PROJECT_CODE_PATTERN.fullmatch(project_code):
        return project_name, ""
    return project_name, project_code


def parse_date(value: object, *, epoch: datetime = CALENDAR_WINDOWS_1900) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value, epoch=epoch)
        except (TypeError, ValueError, OverflowError):
            return None
        return parsed.date() if isinstance(parsed, datetime) else parsed
    text = normalize_text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _find_section_rows(ws: Worksheet) -> dict[str, int]:
    sections: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = normalize_text(cell.value)
            aliases = {
                SECTION_BASIC: SECTION_BASIC,
                SECTION_SIGNOFF: SECTION_SIGNOFF,
                SECTION_PROBLEMS: SECTION_PROBLEMS,
                SECTION_PROBLEMS_V04: SECTION_PROBLEMS,
            }
            canonical = aliases.get(value)
            if canonical and canonical not in sections:
                sections[canonical] = cell.row
    return sections


def _normalized_aliases(aliases: dict[str, tuple[str, ...]]) -> dict[str, str]:
    return {
        normalize_label(alias): key
        for key, values in aliases.items()
        for alias in values
    }


def _find_named_cells(
    ws: Worksheet,
    aliases: dict[str, tuple[str, ...]],
    *,
    start_row: int = 1,
    end_row: int | None = None,
) -> dict[str, list[tuple[int, int]]]:
    lookup = _normalized_aliases(aliases)
    found = {key: [] for key in aliases}
    last_row = min(end_row or ws.max_row, ws.max_row)
    for row_number in range(max(start_row, 1), last_row + 1):
        for column in range(1, ws.max_column + 1):
            key = lookup.get(normalize_label(ws.cell(row_number, column).value))
            if key:
                found[key].append((row_number, column))
    return found


def _cell_reference(row_number: int, column: int) -> str:
    return f"{get_column_letter(column)}{row_number}"


def _read_right_hand_value(
    ws: Worksheet,
    field_cell: tuple[int, int],
    all_field_cells: list[tuple[int, int]],
) -> tuple[object, str]:
    row_number, column = field_cell
    next_field_columns = [
        other_column
        for other_row, other_column in all_field_cells
        if other_row == row_number and other_column > column
    ]
    end_column = min(next_field_columns) - 1 if next_field_columns else ws.max_column
    for value_column in range(column + 1, end_column + 1):
        value = ws.cell(row_number, value_column).value
        if normalize_text(value):
            return value, _cell_reference(row_number, value_column)
    return None, _cell_reference(row_number, column)


def _header_matches(
    ws: Worksheet,
    row_number: int,
    aliases: dict[str, tuple[str, ...]],
) -> dict[str, list[int]]:
    lookup = _normalized_aliases(aliases)
    matches = {key: [] for key in aliases}
    for column in range(1, ws.max_column + 1):
        key = lookup.get(normalize_label(ws.cell(row_number, column).value))
        if key:
            matches[key].append(column)
    return matches


def _locate_header_row(
    ws: Worksheet,
    start_row: int,
    end_row: int,
    aliases: dict[str, tuple[str, ...]],
) -> tuple[int | None, dict[str, list[int]]]:
    best_row: int | None = None
    best_matches = {key: [] for key in aliases}
    best_score = 0
    for row_number in range(start_row + 1, min(end_row, ws.max_row + 1)):
        matches = _header_matches(ws, row_number, aliases)
        score = sum(bool(columns) for columns in matches.values())
        if score > best_score:
            best_row = row_number
            best_matches = matches
            best_score = score
    return best_row, best_matches


def _is_v04_candidate(ws: Worksheet) -> bool:
    section_cells = _find_named_cells(ws, V04_SECTION_ALIASES)
    if any(section_cells.values()):
        return True
    field_cells = _find_named_cells(ws, V04_BASIC_FIELD_ALIASES)
    distinctive_fields = {
        "project_identity",
        "absent_reviewers",
        "meeting_date",
        "stage",
    }
    return any(field_cells[key] for key in distinctive_fields)


def _next_header_row(
    ws: Worksheet, start_row: int, required_labels: tuple[str, ...]
) -> int | None:
    for row_number in range(start_row + 1, min(start_row + 7, ws.max_row) + 1):
        row_text = [normalize_text(ws.cell(row_number, column).value) for column in range(1, ws.max_column + 1)]
        if all(any(label in cell for cell in row_text) for label in required_labels):
            return row_number
    return None


def _header_map(ws: Worksheet, row_number: int) -> dict[str, int]:
    return {
        normalize_text(ws.cell(row_number, column).value): column
        for column in range(1, ws.max_column + 1)
        if normalize_text(ws.cell(row_number, column).value)
    }


def _find_column(headers: dict[str, int], label: str) -> int | None:
    for header, column in headers.items():
        if header == label or label in header:
            return column
    return None


def _basic_info(ws: Worksheet, start_row: int, end_row: int) -> dict[str, object]:
    labels = {
        "技术项目+编号",
        "项目阶段",
        "会议日期",
        "技术项目经理",
        "评审结论",
        "评审意见",
    }
    result: dict[str, object] = {}
    for row_number in range(start_row + 1, end_row):
        for column in range(1, ws.max_column + 1):
            label = normalize_text(ws.cell(row_number, column).value)
            if label not in labels:
                continue
            value: object = None
            for value_column in range(column + 1, ws.max_column + 1):
                candidate = ws.cell(row_number, value_column).value
                if normalize_text(candidate):
                    value = candidate
                    break
            result[label] = value
    return result


def _parse_signoffs(
    ws: Worksheet, header_row: int, end_row: int
) -> list[SignoffRecord]:
    headers = _header_map(ws, header_row)
    columns = {
        "role": _find_column(headers, "评审角色"),
        "reviewer": _find_column(headers, "评审人"),
        "attendance": _find_column(headers, "参会状况"),
        "conclusion": _find_column(headers, "会签结果"),
        "basis": _find_column(headers, "评审依据"),
    }
    if any(column is None for column in columns.values()):
        return []

    signoffs: list[SignoffRecord] = []
    for row_number in range(header_row + 1, end_row):
        role = normalize_text(ws.cell(row_number, columns["role"]).value)
        reviewer_value = ws.cell(row_number, columns["reviewer"]).value
        reviewer = normalize_text(reviewer_value)
        if role.startswith("↓") or reviewer.startswith("↓"):
            continue
        attendance = canonical_attendance(ws.cell(row_number, columns["attendance"]).value)
        conclusion_raw = normalize_text(ws.cell(row_number, columns["conclusion"]).value)
        basis = normalize_text(ws.cell(row_number, columns["basis"]).value)
        if not reviewer and not any((attendance, conclusion_raw, basis)):
            continue
        expert_name, proxy_name, reviewer_raw = parse_reviewer(reviewer_value)
        conclusion, overdue = canonical_conclusion(conclusion_raw)
        signoffs.append(
            SignoffRecord(
                role=role,
                reviewer_raw=reviewer_raw,
                expert_name=expert_name,
                proxy_name=proxy_name,
                attendance=attendance,
                conclusion_raw=conclusion_raw,
                conclusion=conclusion,
                overdue=overdue,
                basis=basis,
                row_number=row_number,
            )
        )
    return signoffs


def _parse_problems(ws: Worksheet, header_row: int) -> list[ProblemRecord]:
    headers = _header_map(ws, header_row)
    columns = {
        "number": _find_column(headers, "序号"),
        "reviewers": _find_column(headers, "评审人"),
        "description": _find_column(headers, "问题描述"),
        "action": _find_column(headers, "建议改善措施"),
        "verification": _find_column(headers, "验证方式与通过条件"),
        "progress": _find_column(headers, "进展"),
        "status": _find_column(headers, "问题状态"),
    }
    if any(column is None for column in columns.values()):
        return []

    problems: list[ProblemRecord] = []
    started = False
    for row_number in range(header_row + 1, ws.max_row + 1):
        number = normalize_text(ws.cell(row_number, columns["number"]).value)
        known_values = {
            key: normalize_text(ws.cell(row_number, column).value)
            for key, column in columns.items()
        }
        if not any(known_values.values()):
            if started:
                break
            continue
        if number.startswith("↓"):
            continue
        other_values = [
            known_values[key]
            for key in ("reviewers", "description", "action", "verification", "progress", "status")
        ]
        if not number and not any(other_values):
            if started:
                break
            continue
        if number and not PROBLEM_NUMBER_PATTERN.fullmatch(number) and not any(other_values):
            continue
        started = True
        problems.append(
            ProblemRecord(
                number=number,
                reviewers=split_reviewers(known_values["reviewers"]),
                description=known_values["description"],
                action=known_values["action"],
                verification=known_values["verification"],
                progress=known_values["progress"],
                status=known_values["status"],
                row_number=row_number,
            )
        )
    return problems


def _parse_sheet(ws: Worksheet, sections: dict[str, int]) -> ReviewSession | None:
    signoff_header = _next_header_row(
        ws, sections[SECTION_SIGNOFF], ("评审角色", "评审人", "参会状况", "会签结果")
    )
    problem_header = _next_header_row(
        ws, sections[SECTION_PROBLEMS], ("序号", "评审人", "问题描述")
    )
    if signoff_header is None or problem_header is None:
        return None

    info = _basic_info(ws, sections[SECTION_BASIC], sections[SECTION_SIGNOFF])
    stage = normalize_text(info.get("项目阶段"))
    if not stage:
        return None
    signoffs = _parse_signoffs(ws, signoff_header, sections[SECTION_PROBLEMS])
    if not signoffs:
        return None
    project_name, project_code = parse_project(info.get("技术项目+编号"))
    session = ReviewSession(
        sheet_name=ws.title,
        project_name=project_name,
        project_code=project_code,
        stage=stage,
        meeting_date=parse_date(info.get("会议日期"), epoch=ws.parent.epoch),
        project_manager=normalize_text(info.get("技术项目经理")),
        meeting_conclusion=canonical_meeting_conclusion(info.get("评审结论")),
        meeting_opinion=normalize_text(info.get("评审意见")),
        signoffs=signoffs,
        problems=_parse_problems(ws, problem_header),
    )
    session.issues.extend(validate_session(session))
    return session


def _parse_v04_signoffs(
    ws: Worksheet,
    header_row: int,
    end_row: int,
    columns: dict[str, int],
) -> list[SignoffRecord]:
    signoffs: list[SignoffRecord] = []
    for row_number in range(header_row + 1, end_row):
        role = normalize_text(ws.cell(row_number, columns["role"]).value)
        reviewer_value = ws.cell(row_number, columns["reviewer"]).value
        reviewer = normalize_text(reviewer_value)
        conclusion_raw = normalize_text(ws.cell(row_number, columns["conclusion"]).value)
        opinion_value = ws.cell(row_number, columns["opinion"]).value
        opinion = normalize_text(opinion_value)
        if role.startswith("↓") or reviewer.startswith("填写"):
            continue
        if not reviewer and not any((role, conclusion_raw, opinion)):
            continue
        expert_name, proxy_name, reviewer_raw = parse_reviewer(reviewer_value)
        conclusion, overdue = canonical_conclusion(conclusion_raw)
        if conclusion == "-":
            conclusion = ""
        opinion_sources: list[OpinionSource] = []
        opinion_items = split_opinion_items(opinion_value)
        for item_index, item in enumerate(opinion_items, start=1):
            reference = _cell_reference(row_number, columns["opinion"])
            if len(opinion_items) > 1:
                reference = f"{reference}-第{item_index}项"
            _add_opinion_source(opinion_sources, item, reference, str(opinion_value or ""))
        signoffs.append(
            SignoffRecord(
                role=role,
                reviewer_raw=reviewer_raw,
                expert_name=expert_name,
                proxy_name=proxy_name,
                attendance="",
                conclusion_raw=conclusion_raw,
                conclusion=conclusion,
                overdue=overdue,
                basis=opinion if opinion != "-" else "",
                row_number=row_number,
                opinion_cell=_cell_reference(row_number, columns["opinion"]),
                cell_references={
                    key: _cell_reference(row_number, column)
                    for key, column in columns.items()
                },
                opinion_sources=opinion_sources,
            )
        )
    return signoffs


def _parse_v04_problems(
    ws: Worksheet,
    header_row: int,
    columns: dict[str, int],
    *,
    source_name: str,
) -> tuple[list[ProblemRecord], list[ValidationIssue]]:
    problems: list[ProblemRecord] = []
    issues: list[ValidationIssue] = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        values = {
            key: normalize_text(ws.cell(row_number, column).value)
            for key, column in columns.items()
        }
        if values["number"].startswith("备注说明"):
            break
        if values["number"].startswith("↓") or values["reviewers"].startswith("填写"):
            continue
        if not any(values.values()):
            continue
        if values["number"] and not any(
            values[key]
            for key in (
                "reviewers",
                "description",
                "action",
                "verification",
                "progress",
                "status",
            )
        ):
            continue
        raw_description = ws.cell(row_number, columns["description"]).value
        description_items = split_numbered_items(raw_description)
        if description_items:
            issues.append(
                ValidationIssue(
                    "problem_description_split",
                    f"问题描述一格包含{len(description_items)}项明确序号，已拆分为多条问题事实",
                    "warning",
                    ws.title,
                    row_number=row_number,
                    cell_reference=_cell_reference(row_number, columns["description"]),
                    source_name=source_name,
                )
            )
        else:
            description_items = [values["description"]]
            if re.search(r"[；;]", values["description"]):
                issues.append(
                    ValidationIssue(
                        "problem_description_maybe_multiple",
                        "问题描述含分号但没有明确序号，已按一条问题保留；请后续一行只记录一个问题",
                        "warning",
                        ws.title,
                        row_number=row_number,
                        cell_reference=_cell_reference(row_number, columns["description"]),
                        source_name=source_name,
                    )
                )
        raw_reviewers = split_reviewers(ws.cell(row_number, columns["reviewers"]).value)
        status_raw = values["status"]
        status = canonical_problem_status(status_raw)
        base_references = {
            key: _cell_reference(row_number, column)
            for key, column in columns.items()
        }
        for item_index, description in enumerate(description_items, start=1):
            number = values["number"]
            references = dict(base_references)
            logical_item_index: int | None = None
            if len(description_items) > 1:
                logical_item_index = item_index
                number = f"{number}.{item_index}" if number else ""
                references["description"] = f"{references['description']}-第{item_index}项"
                references["number"] = f"{references['number']}-第{item_index}项"
            problems.append(
                ProblemRecord(
                    number=number,
                    reviewers=list(raw_reviewers),
                    description=description,
                    action=values["action"],
                    verification=values["verification"],
                    progress=values["progress"],
                    status=status,
                    row_number=row_number,
                    cell_references=references,
                    reviewers_raw=list(raw_reviewers),
                    status_raw=status_raw,
                    source_number=values["number"],
                    item_index=logical_item_index,
                )
            )
    return problems, issues


def _proxy_to_original(signoffs: list[SignoffRecord]) -> dict[str, str]:
    candidates: dict[str, set[str]] = {}
    for signoff in signoffs:
        if signoff.proxy_name and signoff.expert_name:
            candidates.setdefault(signoff.proxy_name, set()).add(signoff.expert_name)
    return {
        proxy_name: next(iter(expert_names))
        for proxy_name, expert_names in candidates.items()
        if len(expert_names) == 1
    }


def _normalize_v04_people(
    signoffs: list[SignoffRecord],
    problems: list[ProblemRecord],
    absent_reviewers: list[str],
    *,
    absent_cell: str,
    sheet_name: str,
    source_name: str,
) -> tuple[list[str], list[ValidationIssue]]:
    issues: list[ValidationIssue] = []
    proxy_to_original = _proxy_to_original(signoffs)
    roster = {signoff.expert_name for signoff in signoffs if signoff.expert_name}
    normalized_absent: list[str] = []
    for name in absent_reviewers:
        normalized = proxy_to_original.get(name, name)
        if normalized != name:
            issues.append(
                ValidationIssue(
                    "absent_proxy_normalized",
                    f"缺席名单中的代理人“{name}”已按区块2唯一代理关系归一到原评审人“{normalized}”",
                    "warning",
                    sheet_name,
                    normalized,
                    cell_reference=absent_cell,
                    source_name=source_name,
                )
            )
        if normalized not in normalized_absent:
            normalized_absent.append(normalized)

    for signoff in signoffs:
        has_valid_conclusion = signoff.conclusion in VALID_CONCLUSIONS
        listed_absent = signoff.expert_name in normalized_absent
        if signoff.proxy_name:
            signoff.attendance = "正常"
            if listed_absent:
                issues.append(
                    ValidationIssue(
                        "absent_with_proxy",
                        f"原评审人“{signoff.expert_name}”虽在缺席名单中，但区块2已记录代理人“{signoff.proxy_name}”履职，按出勤处理",
                        "warning",
                        sheet_name,
                        signoff.expert_name,
                        signoff.row_number,
                        signoff.cell_references.get("reviewer"),
                        source_name,
                    )
                )
        elif has_valid_conclusion:
            signoff.attendance = "正常"
            if listed_absent:
                issues.append(
                    ValidationIssue(
                        "absent_with_valid_signoff",
                        f"评审人“{signoff.expert_name}”虽在缺席名单中，但已提交有效会签，按已履职处理",
                        "warning",
                        sheet_name,
                        signoff.expert_name,
                        signoff.row_number,
                        signoff.cell_references.get("conclusion"),
                        source_name,
                    )
                )
        else:
            signoff.attendance = "缺席未改派" if listed_absent else "正常"

    for problem in problems:
        matched: list[str] = []
        unmatched: list[str] = []
        for raw_name in problem.reviewers_raw:
            normalized = proxy_to_original.get(raw_name, raw_name)
            if normalized in roster:
                if normalized not in matched:
                    matched.append(normalized)
            elif raw_name not in unmatched:
                unmatched.append(raw_name)
        problem.reviewers = matched
        problem.unmatched_reviewers = unmatched

    return normalized_absent, issues


def _attach_problem_opinions(
    signoffs: list[SignoffRecord],
    problems: list[ProblemRecord],
) -> None:
    by_expert = {
        signoff.expert_name: signoff
        for signoff in signoffs
        if signoff.expert_name
    }
    for problem in problems:
        reference = problem.cell_references.get("description", "")
        for reviewer in problem.reviewers:
            signoff = by_expert.get(reviewer)
            if signoff:
                _add_opinion_source(
                    signoff.opinion_sources,
                    problem.description,
                    reference,
                    problem.description,
                )


def _parse_v04_sheet(
    ws: Worksheet,
    *,
    source_name: str,
) -> tuple[ReviewSession | None, list[ValidationIssue]]:
    issues: list[ValidationIssue] = []

    def add_issue(
        code: str,
        message: str,
        *,
        cell_reference: str | None = None,
        severity: str = "error",
    ) -> None:
        issues.append(
            ValidationIssue(
                code,
                message,
                severity,
                ws.title,
                cell_reference=cell_reference,
                source_name=source_name,
            )
        )

    section_cells = _find_named_cells(ws, V04_SECTION_ALIASES)
    section_names = {
        "signoff": SECTION_SIGNOFF,
        "problems": SECTION_PROBLEMS_V04,
    }
    for key, name in section_names.items():
        cells = section_cells[key]
        if not cells:
            add_issue("section_structure_missing", f"缺少业务区块：{name}")
        elif len(cells) > 1:
            references = "、".join(_cell_reference(*cell) for cell in cells)
            add_issue(
                "section_structure_duplicate",
                f"业务区块“{name}”重复，系统不会自行选择",
                cell_reference=references,
            )
    if issues:
        return None, issues

    signoff_section_row = section_cells["signoff"][0][0]
    problem_section_row = section_cells["problems"][0][0]
    if signoff_section_row >= problem_section_row:
        add_issue(
            "section_order_invalid",
            "会签区必须位于问题区之前",
            cell_reference=(
                f"{_cell_reference(*section_cells['signoff'][0])}/"
                f"{_cell_reference(*section_cells['problems'][0])}"
            ),
        )
        return None, issues

    field_cells = _find_named_cells(
        ws,
        V04_BASIC_FIELD_ALIASES,
        end_row=signoff_section_row - 1,
    )
    for key, name in V04_BASIC_FIELD_NAMES.items():
        cells = field_cells[key]
        if not cells:
            add_issue("basic_field_missing", f"基础信息区缺少字段“{name}”")
        elif len(cells) > 1:
            references = "、".join(_cell_reference(*cell) for cell in cells)
            add_issue(
                "basic_field_duplicate",
                f"基础信息区字段“{name}”重复，系统不会自行选择",
                cell_reference=references,
            )
    if issues:
        return None, issues

    all_field_cells = [cell for cells in field_cells.values() for cell in cells]
    field_values: dict[str, object] = {}
    field_references: dict[str, str] = {}
    for key, cells in field_cells.items():
        value, reference = _read_right_hand_value(ws, cells[0], all_field_cells)
        field_values[key] = value
        field_references[key] = reference

    signoff_header, signoff_matches = _locate_header_row(
        ws,
        signoff_section_row,
        problem_section_row,
        V04_SIGNOFF_HEADER_ALIASES,
    )
    problem_header, problem_matches = _locate_header_row(
        ws,
        problem_section_row,
        ws.max_row + 1,
        V04_PROBLEM_HEADER_ALIASES,
    )

    def validate_header(
        kind: str,
        row_number: int | None,
        matches: dict[str, list[int]],
        names: dict[str, str],
    ) -> dict[str, int]:
        columns: dict[str, int] = {}
        if row_number is None:
            add_issue(f"{kind}_header_missing", f"未找到{kind}区表头")
            return columns
        for key, name in names.items():
            matched_columns = matches[key]
            if not matched_columns:
                add_issue(
                    f"{kind}_header_missing",
                    f"{kind}区表头缺少“{name}”",
                    cell_reference=f"A{row_number}",
                )
            elif len(matched_columns) > 1:
                references = "、".join(
                    _cell_reference(row_number, column) for column in matched_columns
                )
                add_issue(
                    f"{kind}_header_duplicate",
                    f"{kind}区表头“{name}”重复，系统不会自行选择",
                    cell_reference=references,
                )
            else:
                columns[key] = matched_columns[0]
        return columns

    signoff_columns = validate_header(
        "signoff", signoff_header, signoff_matches, V04_SIGNOFF_HEADER_NAMES
    )
    problem_columns = validate_header(
        "problem", problem_header, problem_matches, V04_PROBLEM_HEADER_NAMES
    )
    if issues:
        return None, issues

    absent_reviewers_raw = normalize_text(field_values["absent_reviewers"])
    absent_reviewers = split_absent_reviewers(field_values["absent_reviewers"])
    signoffs = _parse_v04_signoffs(
        ws,
        signoff_header,
        problem_section_row,
        signoff_columns,
    )
    problems, problem_issues = _parse_v04_problems(
        ws,
        problem_header,
        problem_columns,
        source_name=source_name,
    )
    issues.extend(problem_issues)
    absent_reviewers, people_issues = _normalize_v04_people(
        signoffs,
        problems,
        absent_reviewers,
        absent_cell=field_references["absent_reviewers"],
        sheet_name=ws.title,
        source_name=source_name,
    )
    issues.extend(people_issues)
    _attach_problem_opinions(signoffs, problems)
    raw_project_identity = normalize_text(field_values["project_identity"])
    raw_stage = normalize_text(field_values["stage"])
    raw_meeting_date = field_values["meeting_date"]
    meeting_conclusion = canonical_meeting_conclusion(field_values["meeting_conclusion"])
    meeting_date = parse_date(raw_meeting_date, epoch=ws.parent.epoch)
    has_business_facts = bool(
        raw_project_identity
        or raw_stage
        or normalize_text(raw_meeting_date)
        or meeting_conclusion
        or problems
        or sum(bool(signoff.expert_name) for signoff in signoffs) >= 3
    )
    if not has_business_facts:
        return None, []

    required_values = {
        "project_identity": raw_project_identity,
        "meeting_date": normalize_text(raw_meeting_date),
        "stage": raw_stage,
        "meeting_conclusion": meeting_conclusion,
    }
    for key, value in required_values.items():
        if not value:
            add_issue(
                "basic_field_value_missing",
                f"字段“{V04_BASIC_FIELD_NAMES[key]}”的对应内容为空",
                cell_reference=field_references[key],
            )

    project_name, project_code = parse_v04_project(raw_project_identity)
    if raw_project_identity and not project_code:
        add_issue(
            "project_identity_invalid",
            "“技术项目名和编码”必须按“完整项目名称-项目编码”填写，且项目编码置于末尾",
            cell_reference=field_references["project_identity"],
        )
    if raw_meeting_date and meeting_date is None:
        add_issue(
            "meeting_date",
            "TDR会议日期缺失或格式无法识别",
            cell_reference=field_references["meeting_date"],
        )

    stage = "TDR2" if raw_stage.upper() == "TDR1+TDR2" else raw_stage
    recorded_reviewer_names = {
        signoff.expert_name
        for signoff in signoffs
        if signoff.expert_name
        and (
            signoff.conclusion in VALID_CONCLUSIONS
            or signoff.conclusion_raw in RECORDED_NO_CONCLUSION_MARKERS
        )
    }
    if len(recorded_reviewer_names) < 3:
        add_issue(
            "effective_signoff_minimum",
            (
                f"正式TDR Sheet仅发现{len(recorded_reviewer_names)}名不同评审人的已记录会签结果；"
                "至少需要3名，允许值为Go／Go with Risk／Redirect／-"
            ),
            cell_reference=(
                f"A{signoff_header}:{get_column_letter(max(signoff_columns.values()))}{signoff_header}"
            ),
        )
    if any(issue.severity == "error" for issue in issues):
        return None, issues

    session = ReviewSession(
        sheet_name=ws.title,
        project_name=project_name,
        project_code=project_code,
        stage=stage,
        meeting_date=meeting_date,
        project_manager="",
        meeting_conclusion=meeting_conclusion,
        meeting_opinion="",
        source_name=source_name,
        absent_reviewers_raw=absent_reviewers_raw,
        absent_reviewers=absent_reviewers,
        signoffs=signoffs,
        problems=problems,
        field_references=field_references,
        parser_profile="v0.4",
    )
    if raw_stage.upper() == "TDR1+TDR2":
        issues.append(
            ValidationIssue(
                "stage_legacy_combined",
                "旧阶段值“TDR1+TDR2”已按TDR2读取，请在源表中改为TDR2",
                "warning",
                ws.title,
                cell_reference=field_references["stage"],
                source_name=source_name,
            )
        )
    session.issues.extend(issues)
    session.issues.extend(validate_session(session))
    return session, []


def read_workbook(
    source: str | Path | bytes | BinaryIO,
    *,
    source_name: str = "",
) -> tuple[list[ReviewSession], list[ValidationIssue]]:
    workbook_source: str | Path | BinaryIO
    if isinstance(source, bytes):
        workbook_source = io.BytesIO(source)
    else:
        workbook_source = source

    workbook = load_workbook(
        workbook_source,
        read_only=False,
        data_only=True,
        keep_links=True,
    )
    sessions: list[ReviewSession] = []
    structure_issues: list[ValidationIssue] = []
    try:
        for ws in workbook.worksheets:
            if "示例" in ws.title:
                continue
            sections = _find_section_rows(ws)
            is_legacy_candidate = SECTION_BASIC in sections
            if not is_legacy_candidate and _is_v04_candidate(ws):
                session, parse_issues = _parse_v04_sheet(ws, source_name=source_name)
                if session:
                    if ws.sheet_state != "visible":
                        session.issues.append(
                            ValidationIssue(
                                "hidden_sheet_parsed",
                                "隐藏Sheet符合有效报告规则，已正常解析；无需删除",
                                "warning",
                                ws.title,
                                source_name=source_name,
                            )
                        )
                    sessions.append(session)
                structure_issues.extend(parse_issues)
                continue
            is_tdr_candidate = ws.title.upper().startswith("TDR") or bool(sections)
            if not all(title in sections for title in SECTION_TITLES):
                if is_tdr_candidate:
                    missing = [title for title in SECTION_TITLES if title not in sections]
                    structure_issues.append(
                        ValidationIssue(
                            "section_structure_missing",
                            "缺少业务区块：" + "、".join(missing),
                            "error",
                            ws.title,
                        )
                    )
                continue
            signoff_header = _next_header_row(
                ws, sections[SECTION_SIGNOFF], SIGNOFF_HEADERS[:4]
            )
            problem_header = _next_header_row(
                ws, sections[SECTION_PROBLEMS], PROBLEM_HEADERS[:3]
            )
            if signoff_header is None:
                structure_issues.append(
                    ValidationIssue("signoff_header_missing", "未找到会签列表表头", "error", ws.title)
                )
            else:
                headers = _header_map(ws, signoff_header)
                missing = [label for label in SIGNOFF_HEADERS if _find_column(headers, label) is None]
                if missing:
                    structure_issues.append(
                        ValidationIssue(
                            "signoff_header_missing",
                            "会签列表缺少列：" + "、".join(missing),
                            "error",
                            ws.title,
                            cell_reference=f"A{signoff_header}",
                        )
                    )
            if problem_header is None:
                structure_issues.append(
                    ValidationIssue("problem_header_missing", "未找到问题汇总表表头", "error", ws.title)
                )
            else:
                headers = _header_map(ws, problem_header)
                missing = [label for label in PROBLEM_HEADERS if _find_column(headers, label) is None]
                if missing:
                    structure_issues.append(
                        ValidationIssue(
                            "problem_header_missing",
                            "问题汇总表缺少列：" + "、".join(missing),
                            "error",
                            ws.title,
                            cell_reference=f"A{problem_header}",
                        )
                    )
            if signoff_header is None or problem_header is None or any(
                issue.sheet_name == ws.title
                and issue.code in {"signoff_header_missing", "problem_header_missing"}
                for issue in structure_issues
            ):
                continue
            session = _parse_sheet(ws, sections)
            if session:
                session.source_name = source_name
                if ws.sheet_state != "visible":
                    session.issues.append(
                        ValidationIssue(
                            "hidden_sheet_parsed",
                            "隐藏Sheet符合有效报告规则，已正常解析；无需删除",
                            "warning",
                            ws.title,
                            source_name=source_name,
                        )
                    )
                sessions.append(session)
    finally:
        workbook.close()

    for session in sessions:
        sheet_stages = {value.upper() for value in re.findall(r"TDR[123]", session.sheet_name, re.IGNORECASE)}
        if sheet_stages and session.stage.upper() not in sheet_stages:
            structure_issues.append(
                ValidationIssue(
                    "sheet_name_stage_mismatch",
                    f"Sheet名称标注阶段“{'／'.join(sorted(sheet_stages))}”与字段“评审阶段”的{session.stage}不一致；已按字段内容读取",
                    "warning",
                    session.sheet_name,
                    cell_reference=session.field_references.get("stage"),
                    source_name=source_name,
                )
            )
    if source_name and sessions:
        filename = Path(source_name).name
        filename_stages = {
            value.upper()
            for value in re.findall(r"TDR[123]", filename, re.IGNORECASE)
        }
        actual_stages = {session.stage.upper() for session in sessions}
        if not filename_stages:
            structure_issues.append(
                ValidationIssue(
                    "filename_stage_missing",
                    "文件名未标明所含TDR阶段；已按Sheet内字段读取，请后续按命名规则补充阶段",
                    "warning",
                    source_name=source_name,
                )
            )
        elif filename_stages != actual_stages:
            structure_issues.append(
                ValidationIssue(
                    "filename_stage_mismatch",
                    (
                        f"文件名阶段“{'／'.join(sorted(filename_stages))}”与有效Sheet阶段"
                        f"“{'／'.join(sorted(actual_stages))}”不一致；已按Sheet内字段读取"
                    ),
                    "warning",
                    source_name=source_name,
                )
            )

    issues = structure_issues + [issue for session in sessions for issue in session.issues]
    issues.extend(validate_stage_conflicts(sessions))
    issues.extend(validate_problem_number_conflicts(sessions))
    if not sessions:
        issues.append(
            ValidationIssue(
                code="no_effective_sessions",
                message="未发现可计分的正式TDR场次；示例表和无正式评审事实的空白占位Sheet已跳过",
                severity="error",
            )
        )
    for issue in issues:
        if issue.source_name is None:
            issue.source_name = source_name
    return sessions, issues
