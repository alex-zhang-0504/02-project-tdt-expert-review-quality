from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime, timezone
from hashlib import sha256
from io import BytesIO
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .scoring import decision_payload

from .models import (
    OpinionSource,
    ProblemRecord,
    ReviewSession,
    SignoffRecord,
    ValidationIssue,
    WorkbookAnalysis,
)


SCHEMA_VERSION = "four-dimension-facts-v0.6"
RULE_VERSION = "facts-v0.6"
PACKAGE_KINDS = {"annual_result", "manager_submission"}
EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PAYLOAD_CHUNK_SIZE = 30_000
_HEADER_FILL = PatternFill("solid", fgColor="0A9BF5")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_LINE = Side(style="thin", color="E2E2E2")
_RULE = Side(style="thin", color="B8B8B8")
_ZEBRA_FILL = PatternFill("solid", fgColor="F6F6F6")


@dataclass(slots=True)
class DimensionOnePackage:
    filename: str
    package_kind: str
    batch_id: str
    manager_id: str
    manager_name: str
    revision: int
    rule_version: str
    exported_at: str
    source_name: str
    sessions: list[ReviewSession]
    issues: list[ValidationIssue]
    payload_hash: str
    decisions: dict

    @property
    def project_codes(self) -> list[str]:
        return sorted({session.project_code for session in self.sessions})


def build_dimension_one_workbook(
    analysis: WorkbookAnalysis,
    *,
    package_kind: str,
    batch_id: str,
    manager_id: str = "",
    manager_name: str = "",
    revision: int = 1,
    product_version: str,
    build_id: str,
) -> bytes:
    if package_kind not in PACKAGE_KINDS:
        raise ValueError("不支持的维度1导出类型")
    if any(issue.severity == "error" for issue in analysis.issues):
        raise ValueError("当前分析仍有错误，不能导出维度1结果")
    batch_id = batch_id.strip()
    manager_id = manager_id.strip()
    manager_name = manager_name.strip()
    if not batch_id:
        raise ValueError("请填写年度批次编号")
    if package_kind == "manager_submission":
        if not manager_id or not manager_name:
            raise ValueError("生成项目经理提交表需要填写项目经理编号和姓名")
        source_managers = {
            session.project_manager.strip()
            for session in analysis.sessions
            if session.project_manager.strip()
        }
        if len(source_managers) > 1:
            raise ValueError("当前报告包含多位技术项目经理，不能生成单人提交表")
        if source_managers and manager_name not in source_managers:
            source_name = next(iter(source_managers))
            raise ValueError(
                f"提交人姓名与TDRX中的技术项目经理“{source_name}”不一致"
            )
    if revision < 1:
        raise ValueError("修订号必须大于等于1")

    exported_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    payload = {
        "schema_version": SCHEMA_VERSION,
        "package_kind": package_kind,
        "batch_id": batch_id,
        "manager_id": manager_id,
        "manager_name": manager_name,
        "revision": revision,
        "rule_version": RULE_VERSION,
        "exported_at": exported_at,
        "product_version": product_version,
        "build_id": build_id,
        "source_name": analysis.source_name,
        "sessions": [asdict(session) for session in analysis.sessions],
        "issues": [asdict(issue) for issue in analysis.issues],
        "decisions": decision_payload(analysis.experts),
    }
    payload_json = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
        default=_json_default,
    )
    payload_hash = sha256(payload_json.encode("utf-8")).hexdigest()

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "V0.6四维事实统计表"
    workbook.properties.subject = SCHEMA_VERSION
    workbook.properties.creator = "V0.6四维事实统计版"

    _write_submission_info(
        workbook,
        payload,
        payload_hash,
        analysis,
    )
    _write_project_list(workbook, analysis)
    _write_fact_detail(workbook, analysis)
    _write_statistics(workbook, analysis)
    _write_issues(workbook, analysis)
    _write_payload(workbook, payload_json, payload_hash, payload)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def load_dimension_one_workbook(content: bytes, filename: str) -> DimensionOnePackage:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"“{filename}”不是可读取的维度1提交表：{exc}") from exc
    required = {"_manifest", "_payload"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"“{filename}”缺少系统提交载荷，不能用于年度汇总")

    manifest_sheet = workbook["_manifest"]
    manifest = {
        str(row[0]): row[1]
        for row in manifest_sheet.iter_rows(values_only=True)
        if row and row[0] is not None
    }
    chunks = [
        str(row[1])
        for row in workbook["_payload"].iter_rows(values_only=True)
        if row and len(row) > 1 and row[0] is not None and row[1] is not None
    ]
    payload_json = "".join(chunks)
    expected_hash = str(manifest.get("payload_hash", ""))
    actual_hash = sha256(payload_json.encode("utf-8")).hexdigest()
    if not expected_hash or actual_hash != expected_hash:
        raise ValueError(f"“{filename}”校验失败，文件可能被修改或损坏")
    try:
        payload = json.loads(payload_json)
    except json.JSONDecodeError as exc:
        raise ValueError(f"“{filename}”的系统提交载荷无法解析") from exc
    if payload.get("schema_version") != SCHEMA_VERSION:
        raise ValueError(f"“{filename}”的提交表版本不受支持")
    if payload.get("rule_version") != RULE_VERSION:
        raise ValueError(f"“{filename}”的维度1规则版本不一致")
    if payload.get("package_kind") != "manager_submission":
        raise ValueError(f"“{filename}”不是项目经理提交表")
    if not payload.get("batch_id") or not payload.get("manager_id"):
        raise ValueError(f"“{filename}”缺少年度批次或项目经理编号")

    sessions = [_review_session_from_dict(item) for item in payload.get("sessions", [])]
    issues = [ValidationIssue(**item) for item in payload.get("issues", [])]
    if not sessions:
        raise ValueError(f"“{filename}”不包含可汇总的维度1事实")
    return DimensionOnePackage(
        filename=filename,
        package_kind=str(payload["package_kind"]),
        batch_id=str(payload["batch_id"]),
        manager_id=str(payload["manager_id"]),
        manager_name=str(payload.get("manager_name", "")),
        revision=int(payload.get("revision", 1)),
        rule_version=str(payload["rule_version"]),
        exported_at=str(payload.get("exported_at", "")),
        source_name=str(payload.get("source_name", "")),
        sessions=sessions,
        issues=issues,
        payload_hash=actual_hash,
        decisions=_validated_decisions(sessions, payload.get("decisions", {})),
    )


def _json_default(value: object) -> str:
    if isinstance(value, date):
        return value.isoformat()
    raise TypeError(f"无法序列化类型：{type(value).__name__}")


def _review_session_from_dict(item: dict[str, object]) -> ReviewSession:
    raw_date = item.get("meeting_date")
    return ReviewSession(
        sheet_name=str(item.get("sheet_name", "")),
        project_name=str(item.get("project_name", "")),
        project_code=str(item.get("project_code", "")),
        stage=str(item.get("stage", "")),
        meeting_date=date.fromisoformat(str(raw_date)) if raw_date else None,
        project_manager=str(item.get("project_manager", "")),
        meeting_conclusion=str(item.get("meeting_conclusion", "")),
        meeting_opinion=str(item.get("meeting_opinion", "")),
        source_name=str(item.get("source_name", "")),
        absent_reviewers_raw=str(item.get("absent_reviewers_raw", "")),
        absent_reviewers=list(item.get("absent_reviewers", [])),
        signoffs=[_signoff_from_dict(value) for value in item.get("signoffs", [])],
        problems=[ProblemRecord(**value) for value in item.get("problems", [])],
        issues=[ValidationIssue(**value) for value in item.get("issues", [])],
        field_references=dict(item.get("field_references", {})),
        parser_profile=str(item.get("parser_profile", "legacy")),
    )


def _signoff_from_dict(item: dict[str, object]) -> SignoffRecord:
    values = dict(item)
    values["opinion_sources"] = [
        OpinionSource(**source) for source in values.get("opinion_sources", [])
    ]
    return SignoffRecord(**values)


def _write_submission_info(
    workbook: Workbook,
    payload: dict[str, object],
    payload_hash: str,
    analysis: WorkbookAnalysis,
) -> None:
    sheet = workbook.create_sheet("00_提交信息")
    rows = [
        ("字段", "内容"),
        ("文件用途", "项目经理维度1提交表" if payload["package_kind"] == "manager_submission" else "集中统计维度1年度结果"),
        ("年度批次", payload["batch_id"]),
        ("项目经理编号", payload["manager_id"]),
        ("项目经理姓名", payload["manager_name"]),
        ("修订号", payload["revision"]),
        ("规则版本", payload["rule_version"]),
        ("数据状态", "试算"),
        ("系统版本", payload["product_version"]),
        ("构建标识", payload["build_id"]),
        ("导出时间UTC", payload["exported_at"]),
        ("源批次", analysis.source_name),
        ("项目数", len({session.project_code for session in analysis.sessions})),
        ("场次数", len(analysis.sessions)),
        ("评审人数", len(analysis.experts)),
        ("载荷校验值", payload_hash),
        ("使用说明", "本文件由系统生成，请勿手工修改。最终汇总会重新计算维度1事实统计。"),
    ]
    for row in rows:
        sheet.append(row)
    _format_sheet(sheet, header_rows=1, widths=(22, 72))


def _write_project_list(workbook: Workbook, analysis: WorkbookAnalysis) -> None:
    sheet = workbook.create_sheet("01_项目清单")
    sheet.append(("项目编码", "子任务名称", "技术项目经理", "TDR3日期", "阶段数", "源报告"))
    grouped: dict[str, list[ReviewSession]] = {}
    for session in analysis.sessions:
        grouped.setdefault(session.project_code, []).append(session)
    for project_code in sorted(grouped):
        sessions = grouped[project_code]
        tdr3_dates = [session.meeting_date for session in sessions if session.stage == "TDR3" and session.meeting_date]
        sheet.append((
            project_code,
            sessions[0].project_name,
            sessions[0].project_manager,
            max(tdr3_dates).isoformat() if tdr3_dates else "",
            len({session.stage for session in sessions}),
            "；".join(sorted({session.source_name for session in sessions if session.source_name})),
        ))
    _format_sheet(sheet, widths=(18, 34, 20, 16, 12, 48))


def _write_fact_detail(workbook: Workbook, analysis: WorkbookAnalysis) -> None:
    sheet = workbook.create_sheet("02_场次事实")
    sheet.append(("评审人", "项目编码", "子任务名称", "阶段", "出勤记录", "代理人", "会签原文", "源报告", "工作表", "单元格"))
    opinions = workbook.create_sheet("03_意见证据")
    opinions.append(("评审人", "项目编码", "子任务名称", "阶段", "意见ID", "意见原文", "AI状态", "对策片段", "理由", "是否计入", "疑似待确认", "识别版本", "人工记录", "源报告", "工作表", "单元格", "原始文本"))
    for expert in analysis.experts:
        for fact in expert.sessions:
            sheet.append((expert.expert_name, fact.project_code, fact.project_name, fact.stage,
                          fact.attendance, fact.proxy_name or "", fact.signoff, fact.source_name,
                          fact.sheet_name, json.dumps(fact.cells, ensure_ascii=False)))
            for opinion in fact.opinions:
                included = opinion.ai_status in {"yes", "suspected"} and opinion.included is not False
                opinions.append((expert.expert_name, fact.project_code, fact.project_name, fact.stage,
                    opinion.opinion_id, opinion.text, opinion.ai_status, opinion.excerpt, opinion.reason,
                    "待AI识别" if opinion.ai_status == "pending" else "是" if included else "否",
                    "疑似待确认" if opinion.ai_status == "suspected" and opinion.included is None else "",
                    opinion.rule_version, json.dumps(opinion.audit, ensure_ascii=False),
                    fact.source_name, fact.sheet_name, "、".join(opinion.cells), "\n".join(opinion.raw_texts)))
    _format_sheet(sheet, widths=(18, 18, 32, 12, 18, 16, 20, 36, 24, 40))
    _format_sheet(opinions, widths=(18, 18, 32, 12, 28, 60, 18, 48, 40, 16, 18, 24, 40, 36, 24, 30, 60))


def _write_statistics(workbook: Workbook, analysis: WorkbookAnalysis) -> None:
    keys = ("attendance_rate", "signoff_rate", "opinion_rate", "solution_rate")
    labels = ("出勤率", "会签率", "意见提出率", "含对策意见率")
    for title, stages in (("04_分阶段统计", ("TDR1", "TDR2", "TDR3")), ("05_全部阶段汇总", ("全部阶段",))):
        sheet = workbook.create_sheet(title)
        sheet.append(["评审人"] + [stage + " " + label for stage in stages for label in labels] + ["评审参与度（有效参评场次）", "代理率"])
        for expert in analysis.experts:
            stats = [expert.overall if stage == "全部阶段" else expert.stages[stage] for stage in stages]
            sheet.append([expert.expert_name] + [
                None if value[key] is None else value[key] / 100 for value in stats for key in keys
            ] + [None if expert.overall["unknown"] else expert.overall["attended"],
                 None if expert.overall["proxy_rate"] is None else expert.overall["proxy_rate"] / 100])
            from openpyxl.comments import Comment
            for index, value in enumerate(stats):
                formulas = [
                    f"实参场次÷应参场次×100％＝{value['attended']}÷{value['expected']}×100％；未知出勤{value['unknown']}场",
                    f"已填会签场次÷应参场次×100％＝{value['signed']}÷{value['expected']}×100％",
                    f"意见总条数÷实参场次×100％＝{value['opinions']}÷{value['attended']}×100％",
                    f"含对策意见÷意见总条数×100％＝{value['solutions']}÷{value['opinions']}×100％；待识别{value['pending']}条，疑似{value['suspected']}条",
                ]
                for offset, formula in enumerate(formulas):
                    sheet.cell(sheet.max_row, 2 + index * 4 + offset).comment = Comment(formula, "V0.6")
            sheet.cell(sheet.max_row, sheet.max_column).comment = Comment(
                f"代理场次÷应参场次×100％＝{expert.overall['proxy']}÷{expert.overall['expected']}×100％", "V0.6")
            sheet.cell(sheet.max_row, sheet.max_column - 1).comment = Comment(
                "全部阶段有效参评场次，按项目编码＋阶段去重，代理归原评审人；参与度分由第四模块在完整批次统一计算。", "V0.7")
        _format_sheet(sheet, widths=tuple([18] + [22] * (len(stages) * 4 + 2)))
        for row in sheet.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = "0" if cell.column == sheet.max_column - 1 else "0.##%"


def _write_issues(workbook: Workbook, analysis: WorkbookAnalysis) -> None:
    sheet = workbook.create_sheet("06_异常清单")
    sheet.append(("级别", "代码", "源报告", "Sheet", "单元格", "评审人", "说明"))
    for issue in analysis.issues:
        sheet.append((
            issue.severity,
            issue.code,
            issue.source_name or "",
            issue.sheet_name or "",
            issue.cell_reference or "",
            issue.expert_name or "",
            issue.message,
        ))
    _format_sheet(sheet, widths=(12, 28, 42, 22, 14, 18, 72))


def _write_payload(
    workbook: Workbook,
    payload_json: str,
    payload_hash: str,
    payload: dict[str, object],
) -> None:
    manifest = workbook.create_sheet("_manifest")
    for row in (
        ("schema_version", SCHEMA_VERSION),
        ("package_kind", payload["package_kind"]),
        ("batch_id", payload["batch_id"]),
        ("manager_id", payload["manager_id"]),
        ("revision", payload["revision"]),
        ("rule_version", RULE_VERSION),
        ("payload_hash", payload_hash),
    ):
        manifest.append(row)
    payload_sheet = workbook.create_sheet("_payload")
    for index, start in enumerate(range(0, len(payload_json), _PAYLOAD_CHUNK_SIZE), start=1):
        payload_sheet.append((index, payload_json[start:start + _PAYLOAD_CHUNK_SIZE]))
    manifest.sheet_state = "veryHidden"
    payload_sheet.sheet_state = "veryHidden"


def _format_sheet(
    sheet,
    *,
    header_rows: int = 1,
    widths: tuple[int, ...],
) -> None:
    sheet.freeze_panes = f"A{header_rows + 1}"
    sheet.auto_filter.ref = sheet.dimensions
    for row_index, row in enumerate(sheet.iter_rows(), start=1):
        for column_index, cell in enumerate(row, start=1):
            if cell.data_type == "f":
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(
                right=_RULE if column_index == 1 else _LINE,
                bottom=_LINE,
            )
            if row_index <= header_rows:
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
            elif row_index % 2 == 1:
                cell.fill = _ZEBRA_FILL
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False
def _validated_decisions(sessions, decisions) -> dict:
    from .scoring import build_facts
    from .countermeasures import validate_predictions
    if not isinstance(decisions, dict):
        raise ValueError("意见识别载荷无效")
    facts = {o.opinion_id: o for e in build_facts(sessions) for s in e.sessions for o in s.opinions}
    if set(decisions) != set(facts):
        raise ValueError("提交表意见身份与原始事实不一致")
    for oid, item in decisions.items():
        if not isinstance(item, dict) or (item.get("included") is not None and type(item["included"]) is not bool):
            raise ValueError("人工复核载荷无效")
        if item.get("text") != facts[oid].text or not isinstance(item.get("audit", []), list):
            raise ValueError("提交表意见证据不一致")
        audit = item.get("audit", [])
        if any(not isinstance(a, dict) or not isinstance(a.get("at"), str)
               or type(a.get("to")) is not bool
               or (a.get("from") is not None and type(a["from"]) is not bool) for a in audit):
            raise ValueError("人工复核记录无效")
        status = item.get("ai_status")
        if status != "suspected" and audit:
            raise ValueError("非疑似意见不能含人工复核记录")
        if status == "pending":
            if item.get("included") is not None:
                raise ValueError("未识别意见不能人工计入")
        else:
            validate_predictions([facts[oid]], [{"id": oid, "status": status,
                "excerpt": item.get("excerpt", ""), "reason": item.get("reason", "")}])
            if status != "suspected" and item.get("included") is not None:
                raise ValueError("非疑似意见不能修改计入选择")
        if status != "pending" and item.get("rule_version") != "countermeasure-v0.6":
            raise ValueError("对策识别规则版本不一致")
    return decisions
