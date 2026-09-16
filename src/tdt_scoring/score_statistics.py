"""Module four: calculate scores from raw facts and saved questionnaires."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .scoring import aggregate
from .subjective import DIMENSIONS, require_analysis


STAGE_WEIGHTS = {"TDR1": 4, "TDR2": 2, "TDR3": 4}
RULE_VERSION = "scores-v0.8"
COMPONENTS = (("attendance", "出勤", 5), ("signoff", "会签", 10),
              ("opinion", "意见", 10))


def stage_score(sessions):
    facts = aggregate(sessions)
    expected, attended, opinions = facts["expected"], facts["attended"], facts["opinions"]
    result = {"applicable": bool(expected), "expected": expected, "attended": attended,
              "opinions": opinions, "suspected": facts["suspected"], "components": {},
              "solutions": facts["solutions"], "solution_bonus": None if facts["pending"] else facts["solutions"] * 2,
              "raw_score": None, "opinion_bonus": 0 if not expected else None, "reasons": []}
    if not expected:
        return result
    reasons = result["reasons"]
    attendance = None if facts["unknown"] else attended / expected * 5
    signoff = facts["signed"] / expected * 10
    if facts["unknown"]:
        reasons.append("出勤状态未知")
    opinion = 0.0 if not opinions else None if facts["unknown"] or not attended else min(opinions / attended, 1) * 10
    if not facts["unknown"] and attended:
        result["opinion_bonus"] = max(opinions - attended, 0)
    elif not facts["unknown"]:
        reasons.append("实参为0，超额意见奖励待确认")
    if opinions and not attended and not facts["unknown"]:
        reasons.append("实参为0但有意见，计分口径待确认")
    if facts["pending"]:
        reasons.append(f"{facts['pending']}条意见待识别对策")
    result["components"] = dict(zip((c[0] for c in COMPONENTS), (attendance, signoff, opinion)))
    if all(value is not None for value in result["components"].values()):
        result["raw_score"] = sum(result["components"].values())
    return result


def questionnaire_score(review):
    items, complete = [], bool(review.get("evaluator", "").strip())
    ratings = review.get("ratings", {})
    for dimension in DIMENSIONS:
        rating = ratings.get(dimension["id"], {})
        option = next((o for o in dimension["options"] if o["id"] == rating.get("option")), None)
        required = rating.get("option") == ("high" if dimension["id"] == "contribution" else "low")
        evidence_missing = bool(option and required and not (rating.get("note", "").strip() and rating.get("project_code", "")))
        complete &= bool(option and not evidence_missing)
        items.append({"dimension": dimension["title"], "option": option["title"] if option else "待评价",
                      "score": option["score"] if option else None, "evidence_missing": evidence_missing})
    return items, sum(item["score"] for item in items) if complete else None


def build_statistics(analysis, scope_confirmed=False):
    require_analysis(analysis)
    facts_by_name = {e.expert_name: aggregate(e.sessions) for e in analysis.experts}
    participation_ready = scope_confirmed and not any(f["unknown"] for f in facts_by_name.values())
    tiers = sorted({f["attended"] for f in facts_by_name.values() if f["attended"] >= 3}, reverse=True)
    rows = []
    for expert in analysis.experts:
        stages = {stage: stage_score([s for s in expert.sessions if s.stage == stage]) for stage in STAGE_WEIGHTS}
        denominator = sum(STAGE_WEIGHTS[stage] for stage, result in stages.items() if result["applicable"])
        facts = facts_by_name[expert.expert_name]
        participation = None
        tier = None
        if participation_ready:
            tier = tiers.index(facts["attended"]) + 1 if facts["attended"] >= 3 else None
            participation = 5 if tier == 1 else 3 if tier == 2 else 0
        reasons, process, objective, bonus, solution_bonus = [], None, None, None, None
        if not scope_confirmed:
            reasons.append("待确认报告范围")
        if not denominator:
            reasons.append("没有适用评审阶段")
        if scope_confirmed and not participation_ready:
            reasons.append("批次存在未知出勤，评审参与度待统计")
        ready = bool(denominator) and all(not r["applicable"] or r["raw_score"] is not None for r in stages.values())
        if ready and scope_confirmed:
            process = sum(r["raw_score"] * STAGE_WEIGHTS[stage] / denominator
                            for stage, r in stages.items() if r["applicable"])
            if participation is not None:
                objective = process + participation
        if scope_confirmed and denominator and all(r["opinion_bonus"] is not None for r in stages.values()):
            bonus = sum(r["opinion_bonus"] for r in stages.values())
        if scope_confirmed and denominator and all(r["solution_bonus"] is not None for r in stages.values()):
            solution_bonus = sum(r["solution_bonus"] for r in stages.values())
        for stage, result in stages.items():
            weight = STAGE_WEIGHTS[stage] / denominator if result["applicable"] else 0
            reasons.extend(f"{stage}：{reason}" for reason in result["reasons"])
            raw_score = result.pop("raw_score")
            result.update(weight=round(weight * 100, 2), score=None if raw_score is None else round(raw_score, 2),
                          contribution=None if raw_score is None else round(raw_score * weight, 2))
            result["components"] = {key: None if value is None else round(value, 2) for key, value in result["components"].items()}
        review = analysis.subjective_reviews.get(expert.expert_name, {})
        items, subjective = questionnaire_score(review)
        if subjective is None:
            reasons.append("主观问卷未完成或必填依据未齐")
        suspected = sum(r["suspected"] for r in stages.values())
        if suspected:
            reasons.append(f"含{suspected}条待确认对策，疑似尚未计入")
        uncapped = None if any(v is None for v in (objective, subjective, bonus, solution_bonus)) else objective + subjective + bonus + solution_bonus
        if uncapped is not None and uncapped > 100:
            reasons.append(f"合计{uncapped:.2f}分，按100分封顶")
        rows.append({"expert_name": expert.expert_name, "projects": len({s.project_code for s in expert.sessions}),
                     "sessions": len(expert.sessions), "stages": stages,
                     "process_total": None if process is None else round(process, 2),
                     "participation_count": None if facts["unknown"] else facts["attended"],
                     "participation_tier": tier, "participation_score": participation,
                     "objective_total": None if objective is None else round(objective, 2), "opinion_bonus": bonus,
                     "solution_bonus": solution_bonus,
                     "subjective_items": items, "subjective_total": subjective, "evaluator": review.get("evaluator", ""),
                     "uncapped_total": None if uncapped is None else round(uncapped, 2),
                     "total": None if uncapped is None else round(min(uncapped, 100), 2),
                     "reasons": reasons, "suspected": suspected})
    return {"rule_version": RULE_VERSION, "scope_confirmed": scope_confirmed, "rows": rows}


def build_statistics_workbook(analysis, scope_confirmed=False):
    data = build_statistics(analysis, scope_confirmed)
    wb = Workbook()
    summary = wb.active
    summary.title = "分数统计（试算）"
    summary.append(["评审人", "项目数", "应参场次", "TDR1（25）", "TDR2（25）", "TDR3（25）",
                    "阶段加权（25）", "总参与评审场次", "参与度数量档", "评审参与度（5）",
                    "评审过程表现（30）", "专业价值贡献（70）", "超额意见奖励", "对策奖励", "封顶前合计", "总分（100）", "状态与说明"])
    stage_sheet = wb.create_sheet("阶段计分依据")
    stage_sheet.append(["评审人", "阶段", "适用", "应参", "实参", "意见条数", "出勤分", "会签分", "意见基础分", "阶段分", "权重％", "加权贡献", "超额意见奖励（不加权）", "含对策意见条数", "对策奖励（不加权）", "说明"])
    questionnaire = wb.create_sheet("主观计分依据")
    questionnaire.append(["评审人", "评价人", "维度", "选项", "分值", "依据状态"])
    for row in data["rows"]:
        summary.append([row["expert_name"], row["projects"], row["sessions"],
                        *[row["stages"][stage]["score"] for stage in STAGE_WEIGHTS],
                        row["process_total"], row["participation_count"], row["participation_tier"], row["participation_score"],
                        row["objective_total"], row["subjective_total"], row["opinion_bonus"], row["solution_bonus"], row["uncapped_total"], row["total"],
                        "；".join(row["reasons"]) or "试算"])
        for stage, result in row["stages"].items():
            stage_sheet.append([row["expert_name"], stage, "适用" if result["applicable"] else "不适用",
                                result["expected"], result["attended"], result["opinions"],
                                *[result["components"].get(c[0]) for c in COMPONENTS], result["score"],
                                result["weight"], result["contribution"], result["opinion_bonus"], result["solutions"], result["solution_bonus"], "；".join(result["reasons"])])
        for item in row["subjective_items"]:
            questionnaire.append([row["expert_name"], row["evaluator"], item["dimension"], item["option"],
                                  item["score"], "待补依据" if item["evidence_missing"] else ""])
    info = wb.create_sheet("使用说明")
    info.append(["规则", RULE_VERSION])
    info.append(["范围确认", "已确认" if scope_confirmed else "未确认，客观合计与总分留空"])
    info.append(["计分", "阶段内出勤5／会签10／意见10分，TDR1／2／3权重4／2／4，不适用阶段权重归一；过程25＋参与度5＋专业价值70。"])
    info.append(["奖励", "各阶段max（意见条数－实参场次，0）直接相加；含对策意见每条另加2分，两类奖励均不乘阶段权重；两维加奖励后最高100分，缺失数据不按0处理。"])
    info.append(["参与度", "本批次有效参评至少3场；前两个不同数量档得5／3分，其余0分，并列同分；全体有效场次确认后统一计算。"])
    info.append(["用途", "仅为试算留档；无自动回载、多人主观合并或排名。未知、待识别与不适用需按状态区分。"])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="0A9BF5")
            cell.font = Font(bold=True, color="FFFFFF")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
