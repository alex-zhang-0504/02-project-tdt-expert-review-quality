"""Six-dimension human assessment, separate from V0.6 fact statistics."""
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field


RULE_VERSION = "subjective-v0.8"
DIMENSIONS = [
    {"id": "preparation", "title": "项目理解与评审准备", "options": [
        {"id": "high", "title": "抓住重点", "score": 10, "description": "理解项目目标、方案与关键约束，能结合材料抓住评审重点。"},
        {"id": "medium", "title": "准备到位", "score": 7, "description": "了解项目基本情况，能围绕本专业完成评审。"},
        {"id": "low", "title": "准备不足", "score": 0, "description": "明显未了解必要材料，反复偏离议题或影响有效评审。"}]},
    {"id": "judgment", "title": "风险识别与专业判断", "options": [
        {"id": "high", "title": "判断深入", "score": 15, "description": "识别关键风险，说明依据、影响及优先级，判断有助于评审决策。"},
        {"id": "medium", "title": "判断合理", "score": 10, "description": "能识别常见问题，判断基本合理，但风险分析不够深入。"},
        {"id": "low", "title": "判断失当", "score": 0, "description": "对职责范围内已有充分线索的关键风险明显漏判，或无依据作出重要判断。"}]},
    {"id": "guidance", "title": "改善建议与方案指导", "options": [
        {"id": "high", "title": "建议可落地", "score": 15, "description": "提出具体可执行的建议，结合项目约束说明方案取舍，帮助推进问题解决。"},
        {"id": "medium", "title": "方向合理", "score": 10, "description": "建议方向合理，但具体措施或适用条件仍需进一步明确。"},
        {"id": "low", "title": "指导不足", "score": 0, "description": "对需要指导的问题只作泛化评价，或提出明显不可执行的建议。"}]},
    {"id": "verification", "title": "验证把关与闭环质量", "options": [
        {"id": "high", "title": "把关有效", "score": 10, "description": "指出关键验证证据及通过条件，复核时识别证据缺口，推动问题有效关闭。"},
        {"id": "medium", "title": "复核到位", "score": 7, "description": "能检查主要验证结果，对明显未满足要求的问题提出补充要求。"},
        {"id": "low", "title": "把关不足", "score": 0, "description": "对负责复核的问题未核实关键证据便认可关闭，或无依据反复变更要求。"}]},
    {"id": "collaboration", "title": "沟通协作与评审担当", "options": [
        {"id": "high", "title": "有据有担当", "score": 10, "description": "观点清楚、依据充分，能协调分歧；关键问题敢于坚持，也能根据新证据修正判断。"},
        {"id": "medium", "title": "沟通尽责", "score": 7, "description": "表达清楚、配合讨论，能够说明并承担本人的专业判断。"},
        {"id": "low", "title": "协作失当", "score": 0, "description": "回避应作出的判断，或以情绪化、无依据的表达妨碍有效讨论。"}]},
    {"id": "contribution", "title": "突出贡献", "options": [
        {"id": "high", "title": "有突出贡献", "score": 10, "description": "有可核实的超出常规履职的贡献：避免重大风险、突破关键难题或推动方案明显改善；普通建议被采纳不足以获得本项加分。"},
        {"id": "low", "title": "无突出贡献", "score": 0, "description": "本期无符合上述标准的可指认贡献，本选项不代表日常履职不合格。"}]},
]


class RatingInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    option: str
    project_code: str = Field(default="", max_length=200)
    note: str = Field(default="", max_length=100)


class ReviewInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    analysis_id: str = Field(min_length=1)
    expert_name: str = Field(min_length=1)
    evaluator: str = Field(min_length=1, max_length=80)
    ratings: dict[str, RatingInput]


def require_analysis(analysis):
    if any(issue.severity == "error" for issue in analysis.issues):
        raise ValueError("请先处理报告中的阻断问题")


def save_review(analysis, payload: ReviewInput) -> dict:
    require_analysis(analysis)
    expert = next((e for e in analysis.experts if e.expert_name == payload.expert_name), None)
    if expert is None:
        raise ValueError("评审人不属于当前分析")
    if not payload.evaluator.strip():
        raise ValueError("请填写评价人")
    catalog = {d["id"]: d for d in DIMENSIONS}
    if payload.ratings.keys() - catalog.keys():
        raise ValueError("存在未知评分维度")
    projects = {s.project_code for s in expert.sessions}
    evidence_missing = False
    ratings = {}
    for dimension_id, rating in payload.ratings.items():
        option = next((o for o in catalog[dimension_id]["options"] if o["id"] == rating.option), None)
        if option is None:
            raise ValueError("该维度不存在所选档位")
        project, note = rating.project_code.strip(), rating.note.strip()
        if project and project not in projects:
            raise ValueError("关联项目不属于该评审人的当前参评范围")
        required = (dimension_id == "contribution" and rating.option == "high") or (
            dimension_id != "contribution" and rating.option == "low")
        evidence_missing |= required and not (project and note)
        ratings[dimension_id] = {"option": rating.option, "project_code": project, "note": note}
    status = "待评价" if len(ratings) < len(DIMENSIONS) else "待补依据" if evidence_missing else "已完成"
    record = {"rule_version": RULE_VERSION, "evaluator": payload.evaluator.strip(),
              "ratings": ratings, "status": status,
              "updated_at": datetime.now(timezone.utc).isoformat()}
    analysis.subjective_reviews[payload.expert_name] = record
    return record


def build_workbook(analysis) -> bytes:
    require_analysis(analysis)
    wb = Workbook()
    summary = wb.active
    summary.title = "主观问卷"
    summary.append(["评审人", "评价人", "状态", "参评项目范围", "更新时间", "规则版本"])
    detail = wb.create_sheet("选档与依据")
    detail.append(["评审人", "维度", "选项", "项目编码", "事实依据", "条目解释"])
    for expert in analysis.experts:
        review = analysis.subjective_reviews.get(expert.expert_name, {})
        projects = sorted({f"{s.project_name}（{s.project_code}）" for s in expert.sessions})
        summary.append([expert.expert_name, review.get("evaluator", ""), review.get("status", "待评价"),
                        "；".join(projects), review.get("updated_at", ""), RULE_VERSION])
        for dimension in DIMENSIONS:
            rating = review.get("ratings", {}).get(dimension["id"], {})
            option = next((o for o in dimension["options"] if o["id"] == rating.get("option")), {})
            detail.append([expert.expert_name, dimension["title"], option.get("title", "待评价"),
                           rating.get("project_code", ""), rating.get("note", ""),
                           option.get("description", "")])
    info = wb.create_sheet("使用说明")
    info.append(["评价范围", analysis.source_name])
    info.append(["保存用途", "本文件用于主观评价查阅留档，不支持自动合并或回载；未完成评价不产生主观总分。"])
    info.append(["计分位置", "本文件只保存问卷，分数统一在第四模块分数统计计算与导出。"])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="0A9BF5")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22 if column[0].column < 5 else 48
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
