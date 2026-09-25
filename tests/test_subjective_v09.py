from copy import deepcopy
from io import BytesIO
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from tests.test_assessment import sample, payload
from tdt_scoring.assessment import confirm_roster
from tdt_scoring.subjective import RULE_VERSION, EvidenceInput, save_review
from tdt_scoring.score_statistics import build_statistics, build_statistics_workbook
from tdt_scoring.manager_evaluation import export_task, import_task


class SubjectiveV09Tests(unittest.TestCase):
    def test_evidence_required_only_for_adverse_ratings_and_contribution(self):
        from tdt_scoring.subjective import DIMENSIONS, rating_result
        for d in DIMENSIONS:
            for o in d['options']:
                required = (d['id'] in ('judgment', 'guidance', 'verification', 'collaboration') and o['id'] == 'low') or (d['id'] == 'contribution' and o['id'] == 'high')
                result = rating_result(d, {'option': o['id']})
                self.assertEqual(result['evidence_missing'], required, (d['id'], o['id']))
                self.assertEqual(result['responded'], not required)

    def test_document_examples_use_only_valid_votes(self):
        from tdt_scoring.manager_evaluation import average_review
        from tdt_scoring.scoring_policy import load_policy
        from tdt_scoring.subjective import DIMENSIONS
        for dimension, options, expected in [('judgment', ['high', 'medium', 'unable'], 12.5),
                                              ('preparation', ['high', 'high', 'high', 'low'], 7.5)]:
            managers = []
            self.a.manager_reviews = {}
            for index, option in enumerate(options):
                mid = str(index)
                managers.append({'manager_id': mid, 'name': '虚拟经理'+mid})
                managers[-1]['experts'] = {self.name: ['B260001']}
                ratings = {d['id']: {'option': 'low' if d['id']=='contribution' else 'high'} for d in DIMENSIONS}
                ratings[dimension] = {'option': option, 'reason': '观察不足', 'project_code': 'B260001', 'note': '可核实的具体事实'}
                self.a.manager_reviews[mid] = {self.name: {'rule_version': RULE_VERSION, 'evaluator': '虚拟经理'+mid, 'ratings': ratings}}
            with patch('tdt_scoring.manager_evaluation.tasks', return_value={'managers': managers, 'unresolved_reports': []}):
                items, total, progress = average_review(self.a, self.name, load_policy()['parameters'])
            index = next(i for i,d in enumerate(DIMENSIONS) if d['id']==dimension)
            self.assertEqual(expected, items[index]['score'])
            self.assertIsNotNone(total)
            self.assertTrue(progress['final'])

    def setUp(self):
        from tdt_scoring import questionnaire
        import json
        data = deepcopy(questionnaire.load()["parameters"])
        data["scores"]["preparation"]["low"] = 0
        fixture = questionnaire.parse(json.dumps(data, ensure_ascii=False).encode())
        mocked = patch("tdt_scoring.questionnaire.load", return_value=fixture)
        mocked.start(); self.addCleanup(mocked.stop)
        self.a = sample()
        self.name = self.a.experts[0].expert_name
        confirm_roster(self.a, [self.name], 'v09')

    def row(self):
        return build_statistics(self.a, True)['rows'][0]

    def save(self, mid, high=True):
        p = payload(self.a, mid, self.name, high)
        save_review(self.a, p)
        return p

    def test_per_question_missing_does_not_discard_other_answers(self):
        self.save('a')
        p = payload(self.a, 'b', self.name, False)
        p.ratings['judgment'].option = 'unable'
        p.ratings['judgment'].reason = '未直接观察相关判断'
        save_review(self.a, p)
        row = self.row()
        self.assertEqual(15, row['subjective_items'][1]['score'])
        self.assertEqual(1, row['subjective_items'][1]['valid_count'])
        self.assertEqual(8.5, row['subjective_items'][0]['score'])
        self.assertEqual(53, row['subjective_total'])
        self.assertEqual(2, row['subjective_progress']['completed'])

    def test_contribution_maximum_and_incomplete_contribution(self):
        p = payload(self.a, 'a', self.name)
        p.ratings['contribution'].option = 'high'
        save_review(self.a, p)
        self.save('b')
        self.assertIsNone(self.row()['subjective_total'])
        p.ratings['contribution'].evidence = [EvidenceInput(project_code='B260001', note='虚拟技术贡献及验证记录')]
        save_review(self.a, p)
        self.assertEqual(70, self.row()['subjective_total'])
        q = payload(self.a, 'b', self.name)
        q.ratings['contribution'].option = 'high'
        q.ratings['contribution'].evidence = [EvidenceInput(project_code='B260002', note='另一条有效贡献')]
        save_review(self.a, q)
        self.assertEqual(70, self.row()['subjective_total'])

    def test_true_zero_is_not_excluded_and_no_valid_answer_has_no_total(self):
        self.save('a')
        p = payload(self.a, 'b', self.name)
        p.ratings['preparation'].option = 'low'
        p.ratings['preparation'].evidence = [EvidenceInput(project_code='B260002', note='多次明显准备不足的具体事实')]
        save_review(self.a, p)
        self.assertEqual(5, self.row()['subjective_items'][0]['score'])
        for mid in ('a', 'b'):
            q = payload(self.a, mid, self.name)
            q.ratings['preparation'].option = 'no_opportunity'
            q.ratings['preparation'].reason = '本周期无相关观察机会'
            save_review(self.a, q)
        self.assertIsNone(self.row()['subjective_total'])
        self.assertIsNone(self.row()['subjective_items'][0]['score'])
        self.assertEqual(2, self.row()['subjective_progress']['completed'])

    def test_missing_reason_and_legacy_version_cannot_finalize(self):
        p = payload(self.a, 'a', self.name)
        p.ratings['preparation'].option = 'unable'
        save_review(self.a, p)
        self.save('b')
        self.assertEqual(1, self.row()['subjective_progress']['completed'])
        self.assertIsNone(self.row()['subjective_total'])
        p.rule_version = 'subjective-v0.8'
        with self.assertRaises(ValueError): save_review(self.a, p)
        self.a.manager_reviews['local:b'][self.name]['rule_version'] = 'subjective-v0.8'
        self.assertEqual(0, self.row()['subjective_items'][0]['valid_count'])

    def test_multiple_evidence_export_and_task_roundtrip(self):
        p = payload(self.a, 'a', self.name)
        p.ratings['contribution'].option = 'high'
        p.ratings['contribution'].evidence = [EvidenceInput(project_code='B260001', note='虚拟记录一'), EvidenceInput(project_code='B260001', note='虚拟记录二')]
        save_review(self.a, p)
        wb = load_workbook(BytesIO(build_statistics_workbook(self.a, True, 'subjective')))
        rows = list(wb['经理评价依据'].values)
        self.assertTrue(any('虚拟记录一' in str(r) for r in rows))
        self.assertTrue(any('虚拟记录二' in str(r) for r in rows))
        self.assertTrue(any(RULE_VERSION in str(r) for r in rows))
        content = export_task(self.a, 'local:a')
        task = load_workbook(BytesIO(content))
        self.assertEqual('能抓住重点', task['经理问卷']['D2'].value)
        self.assertEqual(3, task['补充依据'].max_row)
        before = deepcopy(self.a.manager_reviews)
        self.assertEqual(0, import_task(self.a, content)['updated'])
        self.assertEqual(before, self.a.manager_reviews)


if __name__ == '__main__':
    unittest.main()
