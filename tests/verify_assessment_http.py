"""Run against a local test service: python -m tests.verify_assessment_http URL."""
from io import BytesIO
import json
import sys
from urllib.request import Request, urlopen
from urllib.parse import urlencode, unquote
from urllib.error import HTTPError
from openpyxl import load_workbook
from tests.workbook_factory import build_v04_workbook


def verify(base):
    def request(path, data=None, raw=False, headers=None):
        h = headers or {}
        if isinstance(data, dict):
            data=json.dumps(data).encode(); h['Content-Type']='application/json'
        with urlopen(Request(base+path, data=data, headers=h), timeout=20) as response:
            if raw:
                assert '截止' in unquote(response.headers['Content-Disposition'])
                assert '日前的评审记录.xlsx' in unquote(response.headers['Content-Disposition'])
            content=response.read()
            return content if raw else json.loads(content)
    body=build_v04_workbook([{'stage':'TDR1'},{'stage':'TDR2'},{'stage':'TDR3'}])
    a=request('/api/import/local?filename=virtual.xlsx',body)
    aid=a['analysis_id']; name=a['experts'][0]['expert_name']
    query='?'+urlencode({'analysis_id':aid,'scope_confirmed':'true'})
    try:
        request('/api/statistics/scores'+query)
        raise AssertionError('Unfiltered scores accessible')
    except HTTPError as e: assert e.code==400
    request('/api/assessment/local-manager',dict(analysis_id=aid,source_name='virtual.xlsx',manager_id='http-manager',name='虚拟经理'))
    receipt=request('/api/assessment/policy')
    mixed='@'+name+'、'+name+'；虚拟未匹配'
    preview=request('/api/assessment/roster',dict(analysis_id=aid,names=[mixed]))
    assert preview['included']==[name] and preview['unmatched']==['虚拟未匹配']
    try:
        request('/api/assessment/roster',dict(analysis_id=aid,names=[mixed],batch_id='http-test',confirm=True,policy_hash='stale'))
        raise AssertionError('Stale scoring configuration accepted')
    except HTTPError as e: assert e.code==400
    a=request('/api/assessment/roster',dict(analysis_id=aid,names=[mixed],confirm=True,policy_hash=receipt['sha256']))
    assert a['assessment']['batch_id']==aid
    assert len(a['experts'])==1 and len(a['assessment']['unmatched'])==1
    for kind in ['annual_result','manager_submission']:
        params=dict(analysis_id=aid,package_kind=kind,batch_id='2026',manager_id='http-manager',manager_name='虚拟经理',revision=1)
        direct=load_workbook(BytesIO(request('/api/export/dimension-one?'+urlencode(params),raw=True)))
        assert direct.sheetnames
    task_query='?'+urlencode({'analysis_id':aid,'manager_id':'local:http-manager'})
    wb=load_workbook(BytesIO(request('/api/assessment/task-export'+task_query,raw=True)))
    for row in wb['经理问卷'].iter_rows(min_row=2): row[3].value='low' if row[1].value=='contribution' else 'high'
    output=BytesIO();wb.save(output)
    result=request('/api/assessment/task-import?'+urlencode({'analysis_id':aid}),output.getvalue())
    assert result['updated']==1
    assert request('/api/assessment/task-import?'+urlencode({'analysis_id':aid}),output.getvalue())['updated']==0
    scores=request('/api/statistics/scores'+query)
    assert scores['rows'][0]['subjective_total']==60
    for dimension, sheet in [('objective','客观评分'),('subjective','主观评分'),('all','分数统计（试算）')]:
        exported=load_workbook(BytesIO(request('/api/statistics/scores/export'+query+'&dimension='+dimension,raw=True)))
        assert exported[sheet].max_row==2
    print(json.dumps({'http_checks':'passed','filtered':1,'subjective':60,'exports':3,'duplicate_import':'no additional vote','build':request('/api/health')['build_id']},ensure_ascii=False))


if __name__=='__main__': verify(sys.argv[1])
