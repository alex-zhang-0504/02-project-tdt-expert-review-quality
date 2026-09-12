from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from tdt_scoring.excel_reader import (
    parse_date,
    parse_project,
    parse_reviewer,
    read_workbook,
    split_numbered_items,
)
from tdt_scoring.scoring import build_facts

from tests.workbook_factory import build_v04_workbook, build_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATHS = list(
    (PROJECT_ROOT / "docs").glob("TDT技术项目_TDRX评审报告（模板v0.1）_*.xlsx")
)
if len(TEMPLATE_PATHS) != 1:
    raise RuntimeError("docs目录中必须且只能有一个当前TDRX评审表模板")
TEMPLATE_PATH = TEMPLATE_PATHS[0]


class ExcelReaderTests(unittest.TestCase):
    def test_formal_mode_skips_example_and_empty_sheets(self) -> None:
        sessions, issues = read_workbook(TEMPLATE_PATH)

        self.assertEqual([], sessions)
        issue = next(issue for issue in issues if issue.code == "no_effective_sessions")
        self.assertEqual("error", issue.severity)

    def test_excel_serial_date_is_supported(self) -> None:
        self.assertEqual("2026-03-18", parse_date(46099).isoformat())

    def test_same_expert_is_averaged_across_three_sessions(self) -> None:
        workbook = build_workbook(
            [
                {"stage": "TDR1", "attendance": "正常", "conclusion": "Go"},
                {"stage": "TDR2", "attendance": "缺席未改派", "conclusion": "Go"},
                {"stage": "TDR3", "attendance": "正常", "conclusion": "Go（逾期）"},
            ]
        )
        sessions, _ = read_workbook(workbook)
        score = build_facts(sessions)[0]

        self.assertEqual([True, False, True], [item.attended for item in score.sessions])
        self.assertEqual(3, score.overall["expected"])
        self.assertEqual(66.67, score.overall["attendance_rate"])

    def test_project_and_proxy_parsing_use_last_parentheses(self) -> None:
        self.assertEqual(
            ("虚拟项目（子方向）", "P001"), parse_project("虚拟项目（子方向）（P001）")
        )
        self.assertEqual(
            ("虚拟专家甲", "虚拟专家乙", "虚拟专家甲（虚拟专家乙）"),
            parse_reviewer("虚拟专家甲（虚拟专家乙）"),
        )
        self.assertEqual(
            ("虚拟专家甲", "虚拟专家乙", "虚拟专家甲（代理：虚拟专家乙）"),
            parse_reviewer("虚拟专家甲（代理：虚拟专家乙）"),
        )

    def test_malformed_tdr_sheet_is_not_silently_skipped(self) -> None:
        source = build_workbook(
            [{"stage": "TDR1", "attendance": "正常", "conclusion": "Go"}]
        )
        workbook = load_workbook(BytesIO(source))
        workbook["TDR1"]["A15"] = "错误的区块标题"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual([], sessions)
        self.assertTrue(any(issue.code == "section_structure_missing" for issue in issues))

    def test_v04_title_and_missing_list_drive_session_facts(self) -> None:
        source = build_v04_workbook(
            [
                {
                    "stage": "TDR2",
                    "absent_reviewers": "虚拟专家乙（测试）",
                    "signoffs": [
                        {"role": "射频", "reviewer": "虚拟专家甲", "conclusion": "Go", "opinion": "需关注NTRA线损风险"},
                        {"role": "测试", "reviewer": "虚拟专家乙", "conclusion": "Redirect", "opinion": "测试发现边界场景数据不足，建议补充最差条件验证。"},
                        {"role": "质量", "reviewer": "虚拟专家丙", "conclusion": "-", "opinion": ""},
                    ],
                }
            ],
            project="线损优化-B250001",
        )

        sessions, issues = read_workbook(source, source_name="virtual-v04.xlsx")

        self.assertFalse([issue for issue in issues if issue.severity == "error"])
        self.assertEqual(1, len(sessions))
        session = sessions[0]
        self.assertEqual(("线损优化", "B250001", "TDR2"), (session.project_name, session.project_code, session.stage))
        self.assertEqual("virtual-v04.xlsx", session.source_name)
        self.assertEqual("正常", session.signoffs[0].attendance)
        self.assertEqual("正常", session.signoffs[1].attendance)
        self.assertTrue(any(issue.code == "absent_with_valid_signoff" for issue in issues))
        self.assertEqual("需关注NTRA线损风险", session.signoffs[0].basis)
        self.assertEqual("D9", session.signoffs[0].opinion_cell)

    def test_v04_dash_counts_as_recorded_reviewer_but_scores_zero(self) -> None:
        source = build_v04_workbook(
            [{"stage": "TDR2", "conclusion": "-"}],
        )

        sessions, issues = read_workbook(source)
        scores = {score.expert_name: score for score in build_facts(sessions)}

        self.assertEqual(1, len(sessions))
        self.assertFalse(scores["虚拟专家甲"].sessions[0].signed)
        self.assertFalse(any(issue.code == "signoff_not_provided" for issue in issues))
        self.assertFalse(any(issue.code == "effective_signoff_minimum" for issue in issues))

    def test_v06_blank_signoff_still_counts_in_expected_roster(self) -> None:
        source = build_v04_workbook(
            [
                {
                    "stage": "TDR2",
                    "signoffs": [
                        {"role": "射频", "reviewer": "虚拟专家甲", "conclusion": "Go"},
                        {"role": "测试", "reviewer": "虚拟专家乙", "conclusion": "-"},
                        {"role": "质量", "reviewer": "虚拟专家丙", "conclusion": ""},
                    ],
                }
            ]
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(1, len(sessions))
        self.assertEqual(3, len(sessions[0].signoffs))
        self.assertFalse(any(issue.severity == "error" for issue in issues))

    def test_v04_a1_is_ignored_and_project_field_accepts_dash_variants(self) -> None:
        source = build_v04_workbook(
            [{"stage": "TDR3", "conclusion": "GO"}],
            project="线损优化—B250001",
        )
        workbook = load_workbook(BytesIO(source))
        workbook["TDR3评审报告"]["A1"] = "这里可以填写任意内容"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual(("线损优化", "B250001", "TDR3"), (sessions[0].project_name, sessions[0].project_code, sessions[0].stage))
        self.assertFalse([issue for issue in issues if issue.severity == "error"])

    def test_v04_multi_project_review_uses_first_numbered_project_once(self) -> None:
        source = build_v04_workbook(
            [{"stage": "TDR3"}],
            project=(
                "1. 印度5.5G通信-印度网络拥塞场景优化-B260084H\n"
                "2. 印度5.5G通信-印度UPI扫码支付性能优化-B260084G\n"
                "3. 印度5.5G通信-通信智慧场景感知-B260084A"
            ),
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(1, len(sessions))
        self.assertEqual("印度5.5G通信-印度网络拥塞场景优化", sessions[0].project_name)
        self.assertEqual("B260084H", sessions[0].project_code)
        self.assertFalse([issue for issue in issues if issue.severity == "error"])
        selected = next(issue for issue in issues if issue.code == "multi_project_first_selected")
        self.assertEqual("info", selected.severity)

    def test_v04_multi_project_review_accepts_lettered_items(self) -> None:
        source = build_v04_workbook(
            [{"stage": "TDR3"}],
            project=(
                "a. 印度5.5G通信-印度网络拥塞场景优化-B260084H\n"
                "b. 印度5.5G通信-印度UPI扫码支付性能优化-B260084G\n"
                "c. 印度5.5G通信-通信智慧场景感知-B260084A"
            ),
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(1, len(sessions))
        self.assertEqual("印度5.5G通信-印度网络拥塞场景优化", sessions[0].project_name)
        self.assertEqual("B260084H", sessions[0].project_code)
        self.assertTrue(any(issue.code == "multi_project_first_selected" for issue in issues))

    def test_v04_multi_project_review_accepts_consistent_unnumbered_lines(self) -> None:
        source = build_v04_workbook(
            [{"stage": "TDR3"}],
            project=(
                "印度5.5G通信-印度网络拥塞场景优化-B260084H\n"
                "印度5.5G通信-印度UPI扫码支付性能优化-B260084G\n"
                "印度5.5G通信-通信智慧场景感知-B260084A"
            ),
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(1, len(sessions))
        self.assertEqual("印度5.5G通信-印度网络拥塞场景优化", sessions[0].project_name)
        self.assertEqual("B260084H", sessions[0].project_code)
        self.assertTrue(any(issue.code == "multi_project_first_selected" for issue in issues))

    def test_v04_single_project_manual_line_break_is_not_split(self) -> None:
        source = build_v04_workbook(
            [{"stage": "TDR3"}],
            project="印度5.5G通信-印度网络拥塞\n场景优化-B260084H",
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(1, len(sessions))
        self.assertEqual("印度5.5G通信-印度网络拥塞 场景优化", sessions[0].project_name)
        self.assertEqual("B260084H", sessions[0].project_code)
        self.assertFalse(any(issue.code == "multi_project_first_selected" for issue in issues))

    def test_v04_duplicate_reviewer_roles_merge_into_one_reviewer(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR3",
                "signoffs": [
                    {"role": "评审主席", "reviewer": "高正立", "conclusion": "Go with Risk", "opinion": "需检查网络调度时延"},
                    {"role": "通信专家", "reviewer": "高正立", "conclusion": "Go with Risk", "opinion": "建议补充拥塞场景验证"},
                    {"role": "功耗专家", "reviewer": "虚拟专家乙", "conclusion": "Go", "opinion": ""},
                    {"role": "测试专家", "reviewer": "虚拟专家丙", "conclusion": "-", "opinion": ""},
                ],
            }]
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(3, len(sessions[0].signoffs))
        reviewer = next(item for item in sessions[0].signoffs if item.expert_name == "高正立")
        self.assertEqual("评审主席／通信专家", reviewer.role)
        self.assertEqual(2, len(reviewer.opinion_sources))
        self.assertFalse(any(issue.code == "reviewer_duplicate" for issue in issues))
        scores = [item for item in build_facts(sessions) if item.expert_name == "高正立"]
        self.assertEqual(1, len(scores))
        self.assertEqual(1, scores[0].overall["expected"])

    def test_v04_duplicate_reviewer_conflicting_valid_conclusions_only_warns(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR3",
                "signoffs": [
                    {"role": "评审主席", "reviewer": "高正立", "conclusion": "Go", "opinion": ""},
                    {"role": "通信专家", "reviewer": "高正立", "conclusion": "Redirect", "opinion": ""},
                    {"role": "功耗专家", "reviewer": "虚拟专家乙", "conclusion": "Go", "opinion": ""},
                    {"role": "测试专家", "reviewer": "虚拟专家丙", "conclusion": "-", "opinion": ""},
                ],
            }]
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(3, len(sessions[0].signoffs))
        conflict = next(
            issue for issue in issues if issue.code == "duplicate_reviewer_conclusion_conflict"
        )
        self.assertEqual("warning", conflict.severity)
        self.assertFalse(any(issue.code == "reviewer_duplicate" for issue in issues))

    def test_v04_duplicate_reviewer_prefers_valid_conclusion_over_dash(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR3",
                "signoffs": [
                    {"role": "评审主席", "reviewer": "高正立", "conclusion": "-", "opinion": ""},
                    {"role": "通信专家", "reviewer": "高正立", "conclusion": "Go", "opinion": ""},
                    {"role": "功耗专家", "reviewer": "虚拟专家乙", "conclusion": "Go", "opinion": ""},
                    {"role": "测试专家", "reviewer": "虚拟专家丙", "conclusion": "-", "opinion": ""},
                ],
            }]
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(3, len(sessions[0].signoffs))
        reviewer = next(item for item in sessions[0].signoffs if item.expert_name == "高正立")
        self.assertEqual("Go", reviewer.conclusion)
        self.assertFalse(any(issue.code == "duplicate_reviewer_conclusion_conflict" for issue in issues))
        scores = [item for item in build_facts(sessions) if item.expert_name == "高正立"]
        self.assertTrue(scores[0].sessions[0].signed)

    def test_v04_footer_notes_are_not_parsed_as_problem_rows(self) -> None:
        source = build_v04_workbook([{"stage": "TDR2"}])
        workbook = load_workbook(BytesIO(source))
        workbook["TDR2评审报告"]["A31"] = "备注说明：评审结果说明"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual([], sessions[0].problems)
        self.assertFalse(any(issue.code.startswith("problem_") for issue in issues))

    def test_v04_pre_numbered_blank_problem_rows_are_placeholders(self) -> None:
        source = build_v04_workbook([{"stage": "TDR2"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR2评审报告"]
        sheet["A26"] = 1
        sheet["A27"] = 2
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual([], sessions[0].problems)
        self.assertFalse(any(issue.code.startswith("problem_") for issue in issues))

    def test_v04_close_status_is_accepted_as_transition_alias(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR2",
                "problems": [{
                    "number": "1",
                    "reviewer": "虚拟专家甲",
                    "description": "边界条件需要复核",
                    "status": "close",
                }],
            }]
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(1, len(sessions[0].problems))
        self.assertEqual("closed", sessions[0].problems[0].status)
        self.assertFalse(any(issue.code == "problem_status_alias" for issue in issues))
        self.assertFalse(any(issue.code == "problem_status_invalid" for issue in issues))

    def test_v04_unmatched_absent_reviewer_is_reported_without_guessing(self) -> None:
        source = build_v04_workbook(
            [{"stage": "TDR2", "absent_reviewers": "虚拟专家丁（测试）"}]
        )

        _, issues = read_workbook(source)

        issue = next(issue for issue in issues if issue.code == "absent_reviewer_unmatched")
        self.assertEqual("warning", issue.severity)
        self.assertEqual("B4", issue.cell_reference)

    def test_v04_absent_reviewer_with_dash_scores_as_absent(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR2",
                "absent_reviewers": "虚拟专家甲",
                "signoffs": [
                    {"reviewer": "虚拟专家甲", "conclusion": "-"},
                    {"reviewer": "虚拟专家乙", "conclusion": "Go"},
                    {"reviewer": "虚拟专家丙", "conclusion": "Redirect"},
                ],
            }]
        )

        sessions, issues = read_workbook(source)

        self.assertEqual("缺席未改派", sessions[0].signoffs[0].attendance)
        self.assertFalse(any(issue.code == "absent_with_valid_signoff" for issue in issues))

    def test_v04_proxy_name_in_absent_list_maps_to_original_reviewer(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR2",
                "absent_reviewers": "虚拟专家丁（代理）",
                "signoffs": [
                    {"reviewer": "虚拟专家甲（虚拟专家丁）", "conclusion": "Go"},
                    {"reviewer": "虚拟专家乙", "conclusion": "Go"},
                    {"reviewer": "虚拟专家丙", "conclusion": "Redirect"},
                ],
            }]
        )

        sessions, issues = read_workbook(source)

        self.assertEqual("虚拟专家丁（代理）", sessions[0].absent_reviewers_raw)
        self.assertEqual(["虚拟专家甲"], sessions[0].absent_reviewers)
        self.assertEqual("正常", sessions[0].signoffs[0].attendance)
        proxy_issues = [issue for issue in issues if issue.code == "absent_proxy_resolved"]
        self.assertEqual(1, len(proxy_issues))
        self.assertEqual("info", proxy_issues[0].severity)

    def test_v04_problem_reviewers_intersect_roster_after_proxy_normalization(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR2",
                "signoffs": [
                    {"reviewer": "虚拟专家甲（虚拟专家丁）", "conclusion": "Go"},
                    {"reviewer": "虚拟专家乙", "conclusion": "Go"},
                    {"reviewer": "虚拟专家丙", "conclusion": "Redirect"},
                ],
                "problems": [{
                    "number": "1",
                    "reviewer": "虚拟专家丁、项目参与者戊",
                    "description": "需确认弱网场景边界数据",
                    "status": "open",
                }],
            }]
        )

        sessions, issues = read_workbook(source)
        problem = sessions[0].problems[0]

        self.assertEqual(["虚拟专家甲"], problem.reviewers)
        self.assertEqual(["虚拟专家丁", "项目参与者戊"], problem.reviewers_raw)
        self.assertEqual(["项目参与者戊"], problem.unmatched_reviewers)
        self.assertFalse(any(issue.code == "problem_reviewer_missing" for issue in issues))
        self.assertFalse(any(issue.code == "reviewer_unmatched" for issue in issues))

    def test_v04_problem_reviewer_separators_are_normalized(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR2",
                "problems": [{
                    "number": "1",
                    "reviewer": "虚拟专家甲；虚拟专家乙／项目参与者戊\n虚拟专家丙",
                    "description": "需确认弱网场景边界数据",
                    "status": "open",
                }],
            }]
        )

        sessions, _ = read_workbook(source)
        problem = sessions[0].problems[0]

        self.assertEqual(["虚拟专家甲", "虚拟专家乙", "虚拟专家丙"], problem.reviewers)
        self.assertEqual(["项目参与者戊"], problem.unmatched_reviewers)

    def test_v04_numbered_problem_cell_splits_and_preserves_item_locations(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR2",
                "problems": [{
                    "number": "1",
                    "reviewer": "虚拟专家甲",
                    "description": "1、需确认弱网边界数据\n2．建议补充高温验证",
                    "status": "open",
                }],
            }]
        )

        sessions, issues = read_workbook(source)

        self.assertEqual(["1", "1"], [item.number for item in sessions[0].problems])
        self.assertEqual(
            ["C26-第1项", "C26-第2项"],
            [item.cell_references["description"] for item in sessions[0].problems],
        )
        split_issue = next(issue for issue in issues if issue.code == "problem_description_split")
        self.assertEqual("info", split_issue.severity)
        self.assertFalse(any(issue.code == "problem_number_invalid" for issue in issues))
        self.assertFalse(any(issue.code == "problem_number_duplicate" for issue in issues))

    def test_v04_punctuation_only_problem_stays_one_item_without_reminder(self) -> None:
        for description in (
            "1.价值KPI和成本分析不一致；",
            "需确认弱网边界；建议补充高温验证",
        ):
            with self.subTest(description=description):
                source = build_v04_workbook(
                    [{
                        "stage": "TDR2",
                        "problems": [{
                            "number": "1",
                            "reviewer": "虚拟专家甲",
                            "description": description,
                            "status": "open",
                        }],
                    }]
                )

                sessions, issues = read_workbook(source)

                self.assertEqual(1, len(sessions[0].problems))
                self.assertFalse(
                    any(issue.code == "problem_description_maybe_multiple" for issue in issues)
                )

    def test_common_number_and_letter_markers_split_as_multiple_items(self) -> None:
        cases = (
            "1.第一项\n2.第二项",
            "1、第一项\n2、第二项",
            "a.第一项\nb.第二项",
            "a、第一项\nb、第二项",
            "A）第一项\nB）第二项",
            "（a）第一项\n（b）第二项",
            "一、第一项\n二、第二项",
            "①第一项\n②第二项",
            "1.第一项；2.第二项",
        )

        for description in cases:
            with self.subTest(description=description):
                self.assertEqual(["第一项", "第二项"], split_numbered_items(description))

    def test_v04_filename_never_controls_or_validates_stage(self) -> None:
        source = build_v04_workbook([{"stage": "TDR2"}])

        for source_name in (
            "虚拟项目-B260001.xlsx",
            "虚拟项目-B260001-TDR1.xlsx",
            "虚拟项目-B260001-TDR1+TDR2.xlsx",
            "TDR1能力优化-B260001.xlsx",
        ):
            with self.subTest(source_name=source_name):
                sessions, issues = read_workbook(source, source_name=source_name)

                self.assertEqual("TDR2", sessions[0].stage)
                self.assertFalse(any(issue.code.startswith("filename_stage_") for issue in issues))

    def test_v04_multi_sheet_stages_only_come_from_internal_fields(self) -> None:
        source = build_v04_workbook([
            {"stage": "TDR1"},
            {"stage": "TDR2"},
            {"stage": "TDR3"},
        ])

        sessions, issues = read_workbook(
            source,
            source_name="任意文件名.xlsx",
        )

        self.assertEqual(["TDR1", "TDR2", "TDR3"], [item.stage for item in sessions])
        self.assertFalse(any(issue.code.startswith("filename_stage_") for issue in issues))

    def test_v04_sheet_name_is_not_a_stage_check(self) -> None:
        source = build_v04_workbook([{"stage": "TDR2"}])
        workbook = load_workbook(BytesIO(source))
        workbook["TDR2评审报告"].title = "任意工作表名称"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(
            buffer.getvalue(),
            source_name="虚拟项目-B260001-TDR2.xlsx",
        )

        self.assertEqual("TDR2", sessions[0].stage)
        self.assertFalse(any(issue.code == "sheet_name_stage_mismatch" for issue in issues))

    def test_v04_fields_and_headers_survive_inserted_rows_and_columns(self) -> None:
        source = build_v04_workbook(
            [{"stage": "TDR2", "opinion": "需关注NTRA线损风险"}],
            project="弱网通信体验提升-中高频方案-B250134A",
        )
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR2评审报告"]
        for merged_range in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged_range))
        sheet.insert_rows(1, amount=2)
        sheet.insert_cols(1, amount=2)
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertFalse([issue for issue in issues if issue.severity == "error"])
        self.assertEqual(1, len(sessions))
        session = sessions[0]
        self.assertEqual("弱网通信体验提升-中高频方案", session.project_name)
        self.assertEqual("B250134A", session.project_code)
        self.assertEqual("D5", session.field_references["project_identity"])
        self.assertEqual("F11", session.signoffs[0].opinion_cell)

    def test_v04_signoff_and_problem_sections_survive_internal_blank_rows_and_columns(self) -> None:
        source = build_v04_workbook(
            [{
                "stage": "TDR2",
                "opinion": "需关注NTRA线损风险",
                "problems": [{
                    "number": "1",
                    "reviewer": "虚拟专家甲",
                    "description": "需确认弱网场景边界数据",
                    "status": "open",
                }],
            }],
            project="弱网通信体验提升-中高频方案-B250134A",
        )
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR2评审报告"]
        for merged_range in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged_range))
        sheet.insert_cols(3, amount=2)
        sheet.insert_rows(8, amount=2)
        sheet.insert_rows(27, amount=2)
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertFalse([issue for issue in issues if issue.severity == "error"])
        self.assertEqual(1, len(sessions))
        session = sessions[0]
        self.assertEqual(3, len(session.signoffs))
        self.assertEqual("F11", session.signoffs[0].opinion_cell)
        self.assertEqual(1, len(session.problems))
        self.assertEqual("E30", session.problems[0].cell_references["description"])
        self.assertEqual(["虚拟专家甲"], session.problems[0].reviewers)

    def test_v04_labels_normalize_width_line_breaks_and_spaces(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet["A3"] = "技术项目名\n和编码"
        sheet["A5"] = "ＴＤＲ会议日期"
        sheet["D5"] = "评 审 阶 段"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual(1, len(sessions))
        self.assertFalse([issue for issue in issues if issue.severity == "error"])

    def test_v04_project_stage_field_alias_is_accepted_without_warning(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet["D5"] = "项目阶段"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual(1, len(sessions))
        self.assertEqual("TDR3", sessions[0].stage)
        self.assertFalse([issue for issue in issues if issue.severity != "info"])

    def test_v04_stage_alias_and_canonical_label_together_are_duplicate(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet["F5"] = "项目阶段"
        sheet["G5"] = "TDR3"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual([], sessions)
        issue = next(issue for issue in issues if issue.code == "basic_field_duplicate")
        self.assertIn("评审阶段", issue.message)
        self.assertEqual("D5、F5", issue.cell_reference)

    def test_v04_duplicate_basic_field_blocks_without_guessing(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet["F2"] = "技术项目名和编码"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual([], sessions)
        issue = next(issue for issue in issues if issue.code == "basic_field_duplicate")
        self.assertIn("技术项目名和编码", issue.message)
        self.assertEqual("F2、A3", issue.cell_reference)

    def test_v04_missing_field_name_reports_exact_name(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet["A4"] = "其他名单"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual([], sessions)
        issue = next(issue for issue in issues if issue.code == "basic_field_missing")
        self.assertIn("缺席评审人姓名", issue.message)

    def test_v04_empty_field_value_is_not_read_from_a_fixed_coordinate(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet["B3"] = None
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual([], sessions)
        issue = next(issue for issue in issues if issue.code == "basic_field_value_missing")
        self.assertIn("技术项目名和编码", issue.message)
        self.assertEqual("A3", issue.cell_reference)

    def test_v04_duplicate_header_blocks_without_guessing(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet["E7"] = "评审意见"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual([], sessions)
        issue = next(issue for issue in issues if issue.code == "signoff_header_duplicate")
        self.assertEqual("D7、E7", issue.cell_reference)

    def test_v04_hidden_valid_sheet_is_parsed_without_warning(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet.sheet_state = "hidden"
        workbook.create_sheet("说明")
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual(1, len(sessions))
        self.assertFalse(any(issue.code == "hidden_sheet_parsed" for issue in issues))

    def test_v04_unrelated_auxiliary_sheet_is_silently_skipped(self) -> None:
        source = build_v04_workbook([{"stage": "TDR3"}])
        workbook = load_workbook(BytesIO(source))
        workbook.create_sheet("说明")["A1"] = "项目背景说明"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual(1, len(sessions))
        self.assertFalse(any(issue.sheet_name == "说明" for issue in issues))

    def test_v04_combined_stage_is_invalid_without_normalization(self) -> None:
        source = build_v04_workbook([{"stage": "TDR2"}])
        workbook = load_workbook(BytesIO(source))
        workbook["TDR2评审报告"]["E5"] = "TDR1+TDR2"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual("TDR1+TDR2", sessions[0].stage)
        issue = next(issue for issue in issues if issue.code == "stage_invalid")
        self.assertEqual("error", issue.severity)

    def test_v04_problem_number_format_does_not_trigger_quality_issue(self) -> None:
        for number in ("0", "TDR1-1", "1.1"):
            with self.subTest(number=number):
                source = build_v04_workbook([{
                    "stage": "TDR1",
                    "problems": [{
                        "number": number,
                        "reviewer": "虚拟专家甲",
                        "description": "需复核边界条件",
                        "status": "open",
                    }],
                }])

                _, issues = read_workbook(source)

                self.assertFalse(any(issue.code.startswith("problem_number_") for issue in issues))

    def test_v04_auxiliary_fields_and_headers_do_not_gate_scoring(self) -> None:
        source = build_v04_workbook([{
            "stage": "TDR3",
            "problems": [{
                "reviewer": "虚拟专家甲",
                "description": "需复核弱网边界条件",
            }],
        }])
        workbook = load_workbook(BytesIO(source))
        sheet = workbook["TDR3评审报告"]
        sheet["D3"] = ""
        sheet["E3"] = ""
        sheet["A5"] = ""
        sheet["B5"] = ""
        sheet["A7"] = ""
        for row_number in range(9, 12):
            sheet.cell(row=row_number, column=1, value="")
        for column in (1, 4, 5, 6, 7):
            sheet.cell(row=24, column=column, value="")
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        sessions, issues = read_workbook(buffer.getvalue())

        self.assertEqual(1, len(sessions))
        self.assertEqual("", sessions[0].meeting_conclusion)
        self.assertIsNone(sessions[0].meeting_date)
        self.assertEqual("", sessions[0].signoffs[0].role)
        self.assertEqual(1, len(sessions[0].problems))
        self.assertEqual(["虚拟专家甲"], sessions[0].problems[0].reviewers)
        self.assertFalse([issue for issue in issues if issue.severity == "error"])
        self.assertFalse(any(issue.code.startswith("problem_status_") for issue in issues))

    def test_v04_incomplete_scoring_problem_warns_without_blocking_report(self) -> None:
        source = build_v04_workbook([{
            "stage": "TDR3",
            "problems": [{
                "reviewer": "",
                "description": "需复核弱网边界条件",
                "status": "open",
            }],
        }])

        sessions, issues = read_workbook(source)

        self.assertEqual(1, len(sessions))
        issue = next(issue for issue in issues if issue.code == "problem_reviewer_missing")
        self.assertEqual("warning", issue.severity)
        self.assertFalse([item for item in issues if item.severity == "error"])


if __name__ == "__main__":
    unittest.main()
