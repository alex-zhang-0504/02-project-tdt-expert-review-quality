from dataclasses import asdict
from io import BytesIO
import json
import unittest

from openpyxl import load_workbook
from tdt_scoring.api import app, _encoded
from tdt_scoring.countermeasures import validate_predictions
from tdt_scoring.excel_reader import read_workbook
from tdt_scoring.scoring import build_facts, refresh, decision_payload
from tdt_scoring.service import ScoringService
from tdt_scoring.submission import build_dimension_one_workbook, load_dimension_one_workbook
from tests.workbook_factory import build_workbook, build_v04_workbook


def fixture(stage="TDR1", opinion="1. 请确认温升数据\n2. 建议优化散热结构", **kwargs):
    return build_v04_workbook([{"stage": stage, "opinion": opinion, **kwargs}])


def simulated_classifier(opinions):
    return [{"id": o.opinion_id, "status": "suspected", "excerpt": o.text,
             "reason": "交互测试模拟结果，不代表真实AI准确率"} for o in opinions]


def target(experts):
    return next(e for e in experts if e.expert_name == "虚拟专家甲")


class FactStatisticsTests(unittest.TestCase):
    def test_old_submission_version_is_rejected(self):
        from unittest.mock import patch
        analysis = ScoringService().import_local_bytes(fixture(), "source.xlsx")
        with patch("tdt_scoring.submission.SCHEMA_VERSION", "dimension-one-v0.5"):
            content = build_dimension_one_workbook(analysis, package_kind="manager_submission",
                batch_id="虚拟年度", manager_id="PM01", manager_name="虚拟项目经理",
                revision=1, product_version="v0.5", build_id="test")
        with self.assertRaisesRegex(ValueError, "版本不受支持"):
            load_dimension_one_workbook(content, "old.xlsx")

    def test_invalid_ai_identity_type_is_rejected(self):
        expert = target(build_facts(read_workbook(fixture())[0]))
        with self.assertRaises(ValueError):
            validate_predictions(expert.sessions[0].opinions,
                [{"id": [], "status": "yes", "excerpt": "", "reason": "无效结构"}])

    def test_invalid_audit_and_boolean_payload_are_rejected(self):
        from copy import deepcopy
        from tdt_scoring.submission import _validated_decisions
        service = ScoringService(classifier=simulated_classifier)
        analysis = service.import_local_bytes(fixture(), "source.xlsx")
        service.identify_solutions(analysis.analysis_id)
        original = decision_payload(analysis.experts)
        oid = next(iter(original))
        for invalid in ({"included": 1}, {"audit": ["bad"]}, {"audit": [{"at": "now", "to": "false"}]}):
            with self.subTest(invalid=invalid):
                decisions = deepcopy(original)
                decisions[oid].update(invalid)
                with self.assertRaises(ValueError):
                    _validated_decisions(analysis.sessions, decisions)

    def test_12_sessions_24_opinions_is_200_percent(self):
        sessions = []
        for project in range(4):
            content = build_v04_workbook(
                [{"stage": stage, "opinion": "1. 意见甲\n2. 意见乙"} for stage in ("TDR1", "TDR2", "TDR3")],
                project=f"虚拟项目-B26{project:04d}")
            sessions.extend(read_workbook(content)[0])
        expert = next(e for e in build_facts(sessions) if e.expert_name == "虚拟专家甲")
        self.assertEqual((12, 24, 200), (expert.overall["attended"], expert.overall["opinions"], expert.overall["opinion_rate"]))
        self.assertTrue(all(x["opinion_rate"] == 200 for x in expert.stages.values()))

    def test_three_sessions_nine_opinions_is_300_percent(self):
        content = build_v04_workbook([{"stage": s, "opinion": "1. 甲\n2. 乙\n3. 丙"} for s in ("TDR1", "TDR2", "TDR3")])
        expert = target(build_facts(read_workbook(content)[0]))
        self.assertEqual(300, expert.overall["opinion_rate"])

    def test_overall_recomputes_fraction_not_average_percentages(self):
        sessions = []
        for index, (stage, attendance) in enumerate((("TDR1", "正常"), ("TDR2", "正常"), ("TDR2", "缺席未改派"), ("TDR2", "缺席未改派"))):
            content = build_workbook([{"stage": stage, "project": f"虚拟项目（P{index}）",
                                      "attendance": attendance, "conclusion": "Go"}])
            sessions.extend(read_workbook(content)[0])
        expert = build_facts(sessions)[0]
        self.assertEqual(100, expert.stages["TDR1"]["attendance_rate"])
        self.assertEqual(33.33, expert.stages["TDR2"]["attendance_rate"])
        self.assertEqual(50, expert.overall["attendance_rate"])

    def test_signoff_uses_record_not_conclusion_enum_or_timestamp(self):
        for conclusion, expected in [("", 0), ("-", 0), ("TBD", 100), ("Go（逾期）", 100)]:
            with self.subTest(conclusion=conclusion):
                expert = target(build_facts(read_workbook(fixture(conclusion=conclusion))[0]))
                self.assertEqual(expected, expert.overall["signoff_rate"])

    def test_proxy_stays_with_original_and_has_independent_rate(self):
        content = build_workbook([{"stage": "TDR1", "attendance": "改派（正常）",
            "reviewer": "虚拟专家甲（虚拟代理乙）", "conclusion": "Go", "basis": "建议调整参数"}])
        expert = target(build_facts(read_workbook(content)[0]))
        self.assertEqual("虚拟专家甲", expert.expert_name)
        self.assertEqual(100, expert.overall["proxy_rate"])
        self.assertEqual(100, expert.overall["attendance_rate"])
        self.assertEqual(1, expert.overall["opinions"])

    def test_no_stage_and_unknown_attendance_are_not_zero_rates(self):
        content = build_workbook([{"stage": "TDR1", "attendance": "", "conclusion": "-"}])
        expert = target(build_facts(read_workbook(content)[0]))
        self.assertIsNone(expert.overall["attendance_rate"])
        self.assertEqual(0, expert.stages["TDR3"]["expected"])
        self.assertIsNone(expert.stages["TDR3"]["attendance_rate"])
        self.assertIsNone(expert.overall["solution_rate"])

    def test_identical_opinions_in_different_sessions_are_not_deduped(self):
        content = build_v04_workbook([{"stage": s, "opinion": "建议调整参数"} for s in ("TDR1", "TDR2")])
        expert = target(build_facts(read_workbook(content)[0]))
        self.assertEqual(2, expert.overall["opinions"])
        self.assertEqual(2, len({o.opinion_id for s in expert.sessions for o in s.opinions}))

    def test_missing_ai_does_not_silently_classify(self):
        service = ScoringService()
        analysis = service.import_local_bytes(fixture(), "test.xlsx")
        self.assertEqual("pending", target(analysis.experts).sessions[0].opinions[0].ai_status)
        with self.assertRaisesRegex(ValueError, "尚未配置"):
            service.identify_solutions(analysis.analysis_id)
        self.assertIsNone(target(analysis.experts).overall["solution_rate"])

    def test_suspected_default_included_manual_exclusion_is_audited(self):
        service = ScoringService(classifier=simulated_classifier)
        analysis = service.import_local_bytes(fixture(), "test.xlsx")
        service.identify_solutions(analysis.analysis_id)
        expert = target(analysis.experts)
        self.assertEqual(100, expert.overall["solution_rate"])
        opinion = expert.sessions[0].opinions[0]
        service.select_solution(analysis.analysis_id, opinion.opinion_id, False)
        self.assertEqual(50, expert.overall["solution_rate"])
        self.assertEqual(200, expert.overall["opinion_rate"])
        self.assertEqual(False, opinion.audit[-1]["to"])
        service.select_solution(analysis.analysis_id, opinion.opinion_id, True)
        self.assertEqual(100, expert.overall["solution_rate"])
        self.assertEqual(2, len(opinion.audit))

    def test_untrusted_ai_missing_ids_and_invented_evidence_rejected(self):
        service = ScoringService()
        analysis = service.import_local_bytes(fixture(), "test.xlsx")
        opinions = target(analysis.experts).sessions[0].opinions
        with self.assertRaises(ValueError):
            validate_predictions(opinions, [])
        with self.assertRaises(ValueError):
            validate_predictions(opinions, [{"id": o.opinion_id, "status": "yes", "excerpt": "原文不存在", "reason": "test"} for o in opinions])

    def test_no_old_scores_in_api_or_payload(self):
        analysis = ScoringService().import_local_bytes(fixture(), "test.xlsx")
        encoded = json.dumps(_encoded(analysis))
        for field in ("objective_score", "process_average", "annual_service_score", "total_score", "grade"):
            self.assertNotIn(field, encoded)
        paths = set(app.openapi()["paths"])
        self.assertNotIn("/api/score/finalize", paths)

    def test_manual_decisions_roundtrip_and_recompute_on_merge(self):
        service = ScoringService(classifier=simulated_classifier)
        analysis = service.import_local_bytes(fixture(), "test.xlsx")
        service.identify_solutions(analysis.analysis_id)
        service.select_solution(analysis.analysis_id, target(analysis.experts).sessions[0].opinions[0].opinion_id, False)
        content = build_dimension_one_workbook(analysis, package_kind="manager_submission", batch_id="2026",
            manager_id="PM01", manager_name="虚拟项目经理", product_version="v0.6", build_id="test")
        merged = service.merge_dimension_one_submissions([(content, "submit.xlsx")], expected_manager_count=1)
        self.assertEqual(50, target(merged.experts).overall["solution_rate"])
        self.assertEqual(200, target(merged.experts).overall["opinion_rate"])
        self.assertEqual(1, len(target(merged.experts).sessions[0].opinions[0].audit))
        wb = load_workbook(BytesIO(content))
        self.assertEqual(2, next(row[3] for row in wb["04_分阶段统计"].iter_rows(min_row=2, values_only=True) if row[0] == "虚拟专家甲"))
        self.assertIsNotNone(wb["04_分阶段统计"]["D2"].comment)

    def test_pending_cannot_be_manually_included(self):
        service = ScoringService()
        analysis = service.import_local_bytes(fixture(), "test.xlsx")
        oid = target(analysis.experts).sessions[0].opinions[0].opinion_id
        with self.assertRaises(ValueError):
            service.select_solution(analysis.analysis_id, oid, True)
