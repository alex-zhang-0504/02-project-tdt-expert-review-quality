import json
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from tdt_scoring.assessment import confirm_roster, assign_local_manager, match_roster, tasks, exclude_task, roster_from_excel
from tdt_scoring.manager_evaluation import export_task, import_task
from tdt_scoring.scoring_policy import load_policy, POLICY_PATH
from tdt_scoring.score_statistics import build_statistics, build_statistics_workbook
from tdt_scoring.subjective import DIMENSIONS, ReviewInput, save_review
from tdt_scoring.service import ScoringService
from tests.workbook_factory import build_v04_workbook


def sample():
    service = ScoringService()
    reports = [(build_v04_workbook([{'stage': stage} for stage in stages], project='虚拟项目-'+code), name)
               for code, stages, name in [('B260001', ['TDR1','TDR2'], 'a.xlsx'), ('B260002', ['TDR3'], 'b.xlsx')]]
    analysis = service.import_local_files(reports)
    assign_local_manager(analysis, 'a.xlsx', 'a', '虚拟经理甲')
    assign_local_manager(analysis, 'b.xlsx', 'b', '虚拟经理乙')
    return analysis


def payload(a, mid, name, high=True):
    ratings = {d['id']: {'option': 'high' if high else 'medium'} for d in DIMENSIONS}
    ratings['contribution'] = {'option':'low'}
    return ReviewInput(analysis_id=a.analysis_id, expert_name=name, manager_id='local:'+mid,
                       evaluator='虚拟经理甲' if mid=='a' else '虚拟经理乙', ratings=ratings)


class AssessmentTests(unittest.TestCase):
    def setUp(self):
        self.a = sample()
        self.name = self.a.experts[0].expert_name

    def confirm(self):
        return confirm_roster(self.a, [self.name], 'batch-test')

    def test_intersection_at_name_and_not_fuzzy(self):
        result = match_roster(self.a, ['@'+self.name, self.name, '虚拟未匹配人员'])
        self.assertEqual([self.name], result['included'])
        self.assertEqual(['虚拟未匹配人员'], result['unmatched'])
        self.assertEqual(len(self.a.experts)-1, len(result['excluded']))

    def test_confirmed_roster_fixes_imported_scope_without_extra_checkbox(self):
        self.assertFalse(build_statistics(self.a)['scope_confirmed'])
        self.confirm()
        result = build_statistics(self.a)
        self.assertTrue(result['scope_confirmed'])
        self.assertIsNotNone(result['rows'][0]['objective_total'])
        book = load_workbook(BytesIO(build_statistics_workbook(self.a, dimension='objective')))
        self.assertIn('客观总得分', [c.value for c in book['客观评分'][1]])
        receipt = list(book['客观评分'].values)
        self.assertTrue(any(row[0] == '范围确认' and row[1] is True for row in receipt))
        self.assertEqual(2, sum(row[0] == '纳入报告' for row in receipt))

    def test_horizontal_and_vertical_rosters_match_and_deduplicate(self):
        for separator in ['\n', '\r\n', '、', ',', '，', '.', '。', ';', '；', ' ', '\t', '\u3000']:
            with self.subTest(separator=separator):
                text = separator.join(['@'+self.name, self.name, '虚拟未匹配人员', ''])
                result = match_roster(self.a, [text])
                self.assertEqual([self.name], result['included'])
                self.assertEqual(['虚拟未匹配人员'], result['unmatched'])
                self.assertEqual(result, match_roster(self.a, text))
        self.assertEqual([], match_roster(self.a, ['、， ; \n。'])['requested'])

    def test_changed_policy_receipt_blocks_confirmation(self):
        with self.assertRaisesRegex(ValueError, '评分参数已变化'):
            confirm_roster(self.a, [self.name], 'batch-test', 'stale-hash')
        self.assertFalse(self.a.assessment)
        policy = load_policy()
        result = confirm_roster(self.a, [self.name], 'batch-test', policy['sha256'])
        self.assertEqual(policy['sha256'], result['policy']['sha256'])

    def test_roster_excel_named_column(self):
        wb = Workbook(); wb.active.append(['说明', '姓名']); wb.active.append(['', '@'+self.name])
        self.assertEqual(['@'+self.name], roster_from_excel(self.serialize(wb)))

    def test_feishu_owner_cannot_be_manually_overridden(self):
        self.a.reports[0].source_type = 'feishu_document'
        with self.assertRaises(ValueError): assign_local_manager(self.a, 'a.xlsx', 'other', '虚拟其他经理')

    def test_same_manager_id_cannot_have_two_names(self):
        with self.assertRaises(ValueError): assign_local_manager(self.a, 'b.xlsx', 'a', '虚拟其他经理')

    def test_task_requires_actual_project_intersection(self):
        for expert in self.a.experts:
            if expert.expert_name == self.name:
                for session in expert.sessions:
                    if session.project_code == 'B260002': session.attended = False
        self.confirm()
        self.assertEqual(['local:a'], [m['manager_id'] for m in tasks(self.a)['managers']])

    def test_missing_owner_is_actionable_and_does_not_lock_roster(self):
        self.a.sessions[0].manager_identity = {}
        with self.assertRaisesRegex(ValueError, '项目经理'): self.confirm()
        self.assertFalse(self.a.assessment)

    def test_empty_intersection_rejected_and_preparation_reconfirmation_preserves_policy(self):
        with self.assertRaises(ValueError): confirm_roster(self.a, ['不存在'], 'batch-test')
        self.confirm()
        previous = deepcopy(self.a.assessment['policy'])
        self.confirm()
        self.assertEqual(previous, self.a.assessment['policy'])
        self.a.assessment['completed'] = True
        with self.assertRaises(ValueError): self.confirm()

    def test_policy_snapshot_stable_and_no_silent_fallback(self):
        result = self.confirm()
        expected = deepcopy(result['policy'])
        with patch('tdt_scoring.scoring_policy.POLICY_PATH', Path('missing-config.json')):
            self.assertEqual(expected, build_statistics(self.a, True)['policy'])
            with self.assertRaises(ValueError): load_policy()

    def test_invalid_numbers_and_unknown_dimensions_rejected(self):
        original = json.loads(POLICY_PATH.read_text(encoding='utf-8'))
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'config.json'
            for key, val in [('total_cap', -1), ('precision', True), ('opinion_threshold', 0)]:
                data = deepcopy(original); data[key] = val
                path.write_text(json.dumps(data), encoding='utf-8')
                with patch('tdt_scoring.scoring_policy.POLICY_PATH', path), self.assertRaises(ValueError): load_policy()

    def test_config_changes_score_without_code_changes(self):
        self.confirm()
        before = build_statistics(self.a, True)['rows'][0]
        p = self.a.assessment['policy']['parameters']
        p['components']['attendance'] = 9
        row = build_statistics(self.a, True)['rows'][0]
        self.assertEqual(before['process_total'] + 4, row['process_total'])
        self.assertEqual(before['objective_with_rewards'] + 4, row['objective_with_rewards'])

    def test_two_managers_one_vote_each_despite_three_sessions(self):
        self.confirm()
        catalog = tasks(self.a)
        self.assertEqual(2, len(catalog['managers']))
        save_review(self.a, payload(self.a, 'a', self.name))
        row = build_statistics(self.a, True)['rows'][0]
        self.assertIsNone(row['subjective_total'])
        self.assertEqual(60, row['subjective_progress']['provisional'])
        self.assertEqual(1, row['subjective_progress']['completed'])
        save_review(self.a, payload(self.a, 'b', self.name, False))
        row = build_statistics(self.a, True)['rows'][0]
        self.assertEqual(50.5, row['subjective_total'])
        self.assertEqual(2, row['subjective_progress']['expected'])
        save_review(self.a, payload(self.a, 'a', self.name, False))
        self.assertEqual(41, build_statistics(self.a, True)['rows'][0]['subjective_total'])

    def test_unrelated_manager_or_evidence_rejected(self):
        self.confirm()
        p = payload(self.a, 'a', self.name)
        p.manager_id = 'local:unrelated'
        with self.assertRaises(ValueError): save_review(self.a, p)
        p = payload(self.a, 'a', self.name)
        p.ratings['preparation'].project_code = 'B260002'
        with self.assertRaises(ValueError): save_review(self.a, p)

    def test_incomplete_not_zero_exclusion_and_restore(self):
        self.confirm()
        save_review(self.a, payload(self.a, 'a', self.name))
        p = payload(self.a, 'b', self.name); p.ratings = {}
        save_review(self.a, p)
        self.assertIsNone(build_statistics(self.a, True)['rows'][0]['subjective_total'])
        exclude_task(self.a, 'local:b', self.name, '无足够评价依据')
        self.assertEqual(60, build_statistics(self.a, True)['rows'][0]['subjective_total'])
        with self.assertRaises(ValueError): save_review(self.a, p)
        exclude_task(self.a, 'local:b', self.name, '')
        self.assertIsNone(build_statistics(self.a, True)['rows'][0]['subjective_total'])

    def filled_task(self):
        wb = load_workbook(BytesIO(export_task(self.a, 'local:a')))
        for row in wb['经理问卷'].iter_rows(min_row=2): row[3].value = 'low' if row[1].value=='contribution' else 'high'
        return wb

    def serialize(self, wb):
        output=BytesIO(); wb.save(output); return output.getvalue()

    def test_task_roundtrip_duplicate_and_revision(self):
        self.confirm(); wb = self.filled_task(); content = self.serialize(wb)
        self.assertEqual(1, import_task(self.a, content)['updated'])
        self.assertEqual(0, import_task(self.a, content)['updated'])
        wb['经理问卷']['D2'] = 'medium'
        with self.assertRaises(ValueError): import_task(self.a, self.serialize(wb))
        wb['提交修订']['B1'] = 2
        self.assertEqual(1, import_task(self.a, self.serialize(wb))['updated'])
        with self.assertRaises(ValueError): import_task(self.a, content)
        self.assertEqual(1, len(self.a.manager_reviews['local:a']))

    def test_task_batch_mismatch_and_invalid_option_atomic(self):
        self.confirm(); wb=self.filled_task()
        wb['经理问卷']['D2']='bad'
        with self.assertRaises(ValueError): import_task(self.a, self.serialize(wb))
        self.assertFalse(self.a.manager_reviews)
        wb=self.filled_task(); meta=json.loads(wb['_task']['A1'].value);meta['batch']='other';wb['_task']['A1']=json.dumps(meta)
        with self.assertRaises(ValueError): import_task(self.a, self.serialize(wb))

    def test_independent_exports_and_selected_only(self):
        self.confirm()
        for dim, title in [('objective','客观评分'),('subjective','主观打分')]:
            wb=load_workbook(BytesIO(build_statistics_workbook(self.a, True, dim)))
            if dim == 'subjective':
                self.assertEqual(2, wb[title].max_row)
                self.assertIn('配置及范围', wb.sheetnames)
            if dim == 'objective':
                headers = [cell.value for cell in wb[title][1]]
                self.assertIn('输出有效对策得分', headers)
                self.assertFalse(any('代理' in h for h in headers))
                self.assertEqual(['数据统计', '客观评分'], wb.sheetnames)
                self.assertEqual('A1:K2', wb[title].auto_filter.ref)
                self.assertEqual(5, wb['数据统计'].max_row)
                self.assertEqual('0.0', wb['数据统计']['I2'].number_format)
                self.assertEqual(['全部阶段','TDR1','TDR2','TDR3'], [wb['数据统计'].cell(i,2).value for i in range(2,6)])
                self.assertEqual(self.name, wb['数据统计']['A2'].value)
        row=build_statistics(self.a, True)['rows'][0]
        self.assertIsNotNone(row['objective_with_rewards'])
        self.assertIsNone(row['subjective_total'])


if __name__ == '__main__': unittest.main()
