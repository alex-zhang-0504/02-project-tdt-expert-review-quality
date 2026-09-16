from copy import deepcopy
from io import BytesIO
import unittest

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from tdt_scoring.api import subjective_review
from tdt_scoring.models import ValidationIssue
from tdt_scoring.service import ScoringService
from tdt_scoring.subjective import DIMENSIONS, ReviewInput, build_workbook, save_review
from tdt_scoring.score_statistics import questionnaire_score
from tests.test_scoring import fixture


class SubjectiveTests(unittest.TestCase):
    def setUp(self):
        self.analysis = ScoringService().import_local_bytes(fixture(), "virtual.xlsx")
        self.expert = self.analysis.experts[0]
        self.project = self.expert.sessions[0].project_code

    def payload(self, **changes):
        data = dict(analysis_id=self.analysis.analysis_id, expert_name=self.expert.expert_name,
                    evaluator="虚拟评价人", ratings={d["id"]: {"option": "high"} for d in DIMENSIONS})
        data["ratings"]["contribution"].update(project_code=self.project, note="虚拟案例：提出替代方案避免关键试验失败。")
        data.update(changes)
        return ReviewInput(**data)

    def test_complete_scores_and_facts_are_unchanged(self):
        facts = deepcopy(self.expert.overall)
        review = save_review(self.analysis, self.payload())
        self.assertNotIn("total", review)
        self.assertEqual(70, questionnaire_score(review)[1])
        self.assertEqual("已完成", review["status"])
        payload = self.payload()
        payload.ratings["contribution"].option = "low"
        payload.ratings["contribution"].note = ""
        payload.ratings["contribution"].project_code = ""
        self.assertEqual(60, questionnaire_score(save_review(self.analysis, payload))[1])
        self.assertEqual(facts, self.expert.overall)

    def test_unselected_and_required_evidence_do_not_become_zero_total(self):
        review = save_review(self.analysis, self.payload(ratings={}))
        self.assertIsNone(questionnaire_score(review)[1])
        self.assertEqual("待评价", review["status"])
        payload = self.payload()
        payload.ratings["preparation"].option = "low"
        self.assertEqual("待补依据", save_review(self.analysis, payload)["status"])
        payload.ratings["preparation"].note = "未阅读关键材料，导致讨论反复。"
        self.assertIsNone(questionnaire_score(save_review(self.analysis, payload))[1])
        payload.ratings["preparation"].project_code = self.project
        self.assertEqual(60, questionnaire_score(save_review(self.analysis, payload))[1])

    def test_middle_scores_and_contribution_zero(self):
        ratings = {d["id"]: {"option": "medium" if d["id"] != "contribution" else "low"} for d in DIMENSIONS}
        self.assertEqual(41, questionnaire_score(save_review(self.analysis, self.payload(ratings=ratings)))[1])

    def test_invalid_dimension_option_project_and_person_are_atomic(self):
        saved = deepcopy(save_review(self.analysis, self.payload()))
        cases = [self.payload(ratings={"unknown": {"option": "high"}}),
                 self.payload(ratings={"contribution": {"option": "medium"}}),
                 self.payload(ratings={"preparation": {"option": "low", "project_code": "other"}}),
                 self.payload(expert_name="不存在的虚拟专家"), self.payload(evaluator="   ")]
        for payload in cases:
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                save_review(self.analysis, payload)
        self.assertEqual(saved, self.analysis.subjective_reviews[self.expert.expert_name])

    def test_client_cannot_supply_scores_or_overlong_evidence(self):
        for ratings in [{"preparation": {"option": "high", "score": 99}},
                        {"preparation": {"option": "high", "note": "字" * 101}}]:
            with self.assertRaises(ValidationError):
                self.payload(ratings=ratings)

    def test_blocked_analysis_cannot_save_or_export(self):
        self.analysis.issues.append(ValidationIssue("test", "阻断", "error"))
        with self.assertRaises(ValueError):
            save_review(self.analysis, self.payload())
        with self.assertRaises(ValueError):
            build_workbook(self.analysis)

    def test_missing_analysis_api_returns_404(self):
        with self.assertRaises(HTTPException) as caught:
            subjective_review(self.payload(analysis_id="missing-subjective-analysis"))
        self.assertEqual(404, caught.exception.status_code)

    def test_export_preserves_zero_pending_notes_and_literal_text(self):
        payload = self.payload(evaluator="=1+1")
        save_review(self.analysis, payload)
        wb = load_workbook(BytesIO(build_workbook(self.analysis)))
        row = list(wb["主观问卷"].values)[1]
        self.assertEqual(("=1+1", "已完成"), row[1:3])
        self.assertEqual("s", wb["主观问卷"]["B2"].data_type)
        self.assertTrue(any("替代方案" in str(r[4]) for r in list(wb["选档与依据"].values)[1:]))
        self.assertTrue(any(r[2] == "待评价" for r in list(wb["主观问卷"].values)[2:]))
        for sheet in (wb["主观问卷"], wb["选档与依据"]):
            self.assertFalse(any("分" in str(cell.value) for cell in sheet[1]))
        save_review(self.analysis, self.payload(ratings={}))
        wb = load_workbook(BytesIO(build_workbook(self.analysis)))
        self.assertEqual("待评价", wb["主观问卷"]["C2"].value)
