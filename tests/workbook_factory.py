from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook


def build_workbook(session_rows: list[dict[str, str]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for session in session_rows:
        sheet = workbook.create_sheet(session["stage"])
        rows = [
            ["评审会基本信息列表"],
            ["↓ 填写说明"],
            ["技术项目+编号", session.get("project", "虚拟项目（VIRTUAL-001）")],
            ["项目阶段", session["stage"]],
            ["会议日期", session.get("meeting_date", date(2026, 8, 1))],
            ["技术项目经理", session.get("project_manager", "虚拟项目经理")],
            ["评审结论", session.get("meeting_conclusion", "Go")],
            ["评审意见", ""],
            [],
            ["评审结论和会签列表"],
            ["评审角色", "评审人", "参会状况", "会签结果", "评审依据（数据&标准&测试或技术逻辑等）"],
            ["↓ 填写说明", "填写全名", "下拉", "下拉", "按需填写"],
            [session.get("role", "评审主席"), session.get("reviewer", "虚拟专家甲"), session["attendance"], session["conclusion"], session.get("basis", "")],
            [],
            ["评审问题&建议汇总表"],
            ["序号", "评审人", "问题描述", "建议改善措施", "验证方式与通过条件", "进展（技术负责人跟踪）", "问题状态"],
            ["↓TDRX-No.", "填写人", "填写问题", "填写措施", "填写条件", "填写进展", "open／closed"],
        ]
        if session.get("problem"):
            rows.append(
                [
                    session.get("problem_number", "1"),
                    session.get("problem_reviewer", "虚拟专家甲"),
                    session.get("problem_description", "虚拟问题"),
                    session.get("action", ""),
                    session.get("verification", ""),
                    "",
                    session.get("problem_status", "open"),
                ]
            )
        if session.get("duplicate_reviewer"):
            rows.insert(
                13,
                ["评审专家", "虚拟专家甲", "正常", "Go", ""],
            )
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def build_v04_workbook(
    session_rows: list[dict[str, object]],
    *,
    project: str = "虚拟大TDT-虚拟子任务-B260001",
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for session in session_rows:
        stage = str(session["stage"])
        sheet = workbook.create_sheet(f"{stage}评审报告")
        sheet.append([session.get("title", "第一区块内容不参与关键字段识别")])
        sheet.append(["↓ 填写说明"])
        sheet.append(["技术项目名和编码", project, "", "评审结论", session.get("meeting_conclusion", "Go")])
        sheet.append([
            "缺席评审人姓名",
            session.get("absent_reviewers", "无"),
            "",
            "评审意见",
            session.get("meeting_opinion", ""),
        ])
        sheet.append(["TDR会议日期", session.get("meeting_date", date(2026, 8, 1)), "", "评审阶段", stage])
        sheet.merge_cells("B3:C3")
        sheet.merge_cells("B4:C4")
        sheet.merge_cells("B5:C5")
        sheet.append(["评审结论和会签列表"])
        sheet.append(["评审角色", "评审人姓名", "会签结果", "评审意见"])
        sheet.append(["↓ 填写说明", "填写全名", "未给意见选择[-]", "写清具体技术对象和专业判断"])
        signoffs = session.get("signoffs") or [
            {
                "role": session.get("role", "评审主席"),
                "reviewer": session.get("reviewer", "虚拟专家甲"),
                "conclusion": session.get("conclusion", "Go"),
                "opinion": session.get("opinion", ""),
            },
            {"role": "领域专家", "reviewer": "虚拟专家乙", "conclusion": "-", "opinion": ""},
            {"role": "测试专家", "reviewer": "虚拟专家丙", "conclusion": "-", "opinion": ""},
        ]
        for signoff in signoffs:
            sheet.append([
                signoff.get("role", ""),
                signoff.get("reviewer", ""),
                signoff.get("conclusion", ""),
                signoff.get("opinion", ""),
            ])
        while sheet.max_row < 22:
            sheet.cell(row=sheet.max_row + 1, column=1, value="")
        sheet.append(["评审问题汇总表"])
        sheet.append(["序号", "评审人", "问题描述", "是否要修正", "问题等级", "反馈/修改说明", "问题状态"])
        sheet.append(["↓1", "填写人", "填写问题", "是／否", "一般／严重", "填写反馈", "open／closed"])
        for problem in session.get("problems", []):
            sheet.append([
                problem.get("number", ""),
                problem.get("reviewer", ""),
                problem.get("description", ""),
                problem.get("must_fix", ""),
                problem.get("level", ""),
                problem.get("feedback", ""),
                problem.get("status", ""),
            ])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
