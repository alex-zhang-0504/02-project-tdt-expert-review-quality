from io import BytesIO
import unittest

from openpyxl import load_workbook

from tdt_scoring.api import subjective_catalog
from tdt_scoring.models import ExpertFacts, OpinionFact, SessionFact, WorkbookAnalysis
from tdt_scoring.scoring import aggregate
from tdt_scoring.score_statistics import build_statistics, build_statistics_workbook, stage_score
from tdt_scoring.subjective import DIMENSIONS


def session(stage="TDR1", statuses=(), attended=True, signed=True, project="VIRTUAL"):
    opinions = [OpinionFact(str(i), "虚拟意见", [], [], ai_status=status) for i, status in enumerate(statuses)]
    return SessionFact(project + stage, project, "虚拟项目", stage, stage, "virtual.xlsx",
                       "正常", attended, "Go" if signed else "", signed, None, opinions)


def analysis(sessions):
    expert = ExpertFacts("虚拟评审人", sessions, {}, aggregate(sessions))
    value = WorkbookAnalysis("test", "local", "virtual.xlsx", [], [expert], [])
    value.subjective_reviews[expert.expert_name] = {
        "evaluator": "虚拟评价人", "total": 999, "status": "已完成",
        "ratings": {d["id"]: {"option": "low" if d["id"] == "contribution" else "high"} for d in DIMENSIONS},
    }
    return value


class ScoreStatisticsTests(unittest.TestCase):
    def test_pending_solution_does_not_block_basic_score(self):
        row = self.row([session(statuses=("pending", "no"))])
        self.assertEqual(25, row["objective_total"])
        self.assertEqual(1, row["opinion_bonus"])
        self.assertIsNone(row["solution_bonus"])
        self.assertIsNone(row["total"])

    def test_add_problem_never_reduces_score_and_solution_adds_two(self):
        facts = session(statuses=("yes", "no"))
        before = self.row([facts])
        facts.opinions.append(OpinionFact("extra", "新增问题", [], [], ai_status="no"))
        after_problem = self.row([facts])
        self.assertEqual(before["total"] + 1, after_problem["total"])
        facts.opinions[-1].ai_status = "yes"
        after_solution = self.row([facts])
        self.assertEqual(after_problem["total"] + 2, after_solution["total"])
        self.assertEqual(before["objective_total"], after_solution["objective_total"])

    def test_four_sessions_six_opinions_two_solutions_example(self):
        value = analysis([session(project=str(i), statuses=("yes", "yes", "no", "no", "no", "no") if i == 0 else ()) for i in range(4)])
        value.experts.append(ExpertFacts("虚拟甲", [session(project=f"other{i}") for i in range(6)], {}, {}))
        row = build_statistics(value, True)["rows"][0]
        self.assertEqual((28, 60, 2, 4, 94), (row["objective_total"], row["subjective_total"], row["opinion_bonus"], row["solution_bonus"], row["total"]))

    def test_bonus_below_at_and_above_target(self):
        for count, base, bonus in [(2, 5, 0), (4, 10, 0), (5, 10, 1), (6, 10, 2), (8, 10, 4)]:
            sessions = [session(project=str(i), statuses=("yes",) * count if i == 0 else ()) for i in range(4)]
            result = stage_score(sessions)
            self.assertEqual(base, result["components"]["opinion"])
            self.assertEqual(bonus, result["opinion_bonus"])

    def test_bonus_is_not_weighted_or_offset_by_other_stage(self):
        row = self.row([session(statuses=("yes",) * 5), session("TDR2"), session("TDR3")])
        self.assertEqual(4, row["opinion_bonus"])
        self.assertEqual(19, row["process_total"])
        self.assertEqual(10, row["solution_bonus"])
        self.assertEqual(98, row["total"])

    def test_bonus_caps_total_and_cannot_bypass_missing_questionnaire(self):
        value = analysis([session(statuses=("yes",) * 30)])
        row = build_statistics(value, True)["rows"][0]
        self.assertEqual((29, 174, 100), (row["opinion_bonus"], row["uncapped_total"], row["total"]))
        value.subjective_reviews.clear()
        self.assertIsNone(build_statistics(value, True)["rows"][0]["total"])

    def test_participation_ties_threshold_and_unknown_batch(self):
        value = analysis([])
        value.experts = [ExpertFacts(f"虚拟{i}", [session(project=str(j)) for j in range(count)], {}, {})
                         for i, count in enumerate([6, 6, 4, 3, 2])]
        rows = build_statistics(value, True)["rows"]
        self.assertEqual([5, 5, 3, 0, 0], [r["participation_score"] for r in rows])
        value.experts[-1].sessions[0].attended = None
        self.assertTrue(all(r["participation_score"] is None for r in build_statistics(value, True)["rows"]))

    def test_zero_attendance_bonus_is_pending_and_scope_keeps_bonus_blank(self):
        row = self.row([session(attended=False, signed=False)])
        self.assertIsNone(row["opinion_bonus"])
        self.assertIsNone(row["total"])
        row = self.row([session(statuses=("yes",) * 2)], False)
        self.assertIsNone(row["opinion_bonus"])
        self.assertIsNone(row["participation_score"])

    def row(self, sessions, confirmed=True):
        return build_statistics(analysis(sessions), confirmed)["rows"][0]

    def test_424_weights_and_subjective_calculated_from_choices(self):
        sessions = [session(statuses=("yes", "no")),
                    session("TDR2", ("no",), signed=False), session("TDR3", ("no",))]
        row = self.row(sessions)
        self.assertEqual([25, 15, 25], [s["score"] for s in row["stages"].values()])
        self.assertEqual([40, 20, 40], [s["weight"] for s in row["stages"].values()])
        self.assertEqual((28, 60, 91), (row["objective_total"], row["subjective_total"], row["total"]))

    def test_missing_stages_renormalize_but_absence_remains_applicable(self):
        for stages, weights in [(('TDR1', 'TDR2'), (66.67, 33.33, 0)),
                               (('TDR1', 'TDR3'), (50, 0, 50)), (('TDR2',), (0, 100, 0))]:
            row = self.row([session(s, attended=False, signed=False) for s in stages])
            self.assertEqual(weights, tuple(s["weight"] for s in row["stages"].values()))
            self.assertEqual(0, row["objective_total"])

    def test_scope_unconfirmed_retains_subscores_but_no_total(self):
        row = self.row([session()], False)
        self.assertEqual(15, row["stages"]["TDR1"]["score"])
        self.assertIsNone(row["objective_total"])
        self.assertIsNone(row["total"])

    def test_pending_and_unknown_block_total_without_reweighting(self):
        for bad in [session("TDR3", attended=None),
                    session("TDR3", ("yes",), attended=False)]:
            row = self.row([session("TDR1", ("yes",)), bad])
            self.assertEqual(50, row["stages"]["TDR3"]["weight"])
            self.assertIsNone(row["objective_total"])
            self.assertIsNone(row["total"])
            self.assertTrue(row["stages"]["TDR3"]["reasons"])

    def test_confirmed_zero_solution_is_zero_despite_null_fact_rate(self):
        sessions = [session(statuses=("no",))]
        self.assertIsNone(aggregate(sessions)["solution_rate"])
        self.assertEqual(0, stage_score(sessions)["solution_bonus"])
        self.assertEqual(25, self.row(sessions)["objective_total"])
        self.assertEqual(15, self.row([session()])["objective_total"])

    def test_counts_aggregate_before_ratio_and_cap_does_not_change_facts(self):
        sessions = [session(statuses=("yes", "yes", "yes")), session(project="SECOND")]
        self.assertEqual(150, aggregate(sessions)["opinion_rate"])
        self.assertEqual(25, self.row(sessions)["objective_total"])
        self.assertEqual(150, aggregate(sessions)["opinion_rate"])

    def test_suspected_requires_explicit_inclusion(self):
        value = session(statuses=("suspected",))
        self.assertEqual(0, self.row([value])["solution_bonus"])
        self.assertEqual(1, self.row([value])["suspected"])
        value.opinions[0].included = True
        self.assertEqual(2, self.row([value])["solution_bonus"])
        self.assertEqual(0, self.row([value])["suspected"])
        value.opinions[0].included = False
        self.assertEqual(0, self.row([value])["solution_bonus"])

    def test_incomplete_questionnaire_does_not_use_cached_total(self):
        value = analysis([session()])
        value.subjective_reviews["虚拟评审人"]["ratings"].pop("preparation")
        row = build_statistics(value, True)["rows"][0]
        self.assertIsNone(row["subjective_total"])
        self.assertIsNone(row["total"])

    def test_export_keeps_pending_blank_and_records_scope(self):
        value = analysis([session()])
        wb = load_workbook(BytesIO(build_statistics_workbook(value)))
        self.assertIsNone(wb["分数统计（试算）"]["K2"].value)
        self.assertIsNone(wb["分数统计（试算）"]["P2"].value)
        wb = load_workbook(BytesIO(build_statistics_workbook(value, True)))
        self.assertEqual(75, wb["分数统计（试算）"]["P2"].value)
        self.assertEqual("已确认", wb["使用说明"]["B2"].value)

    def test_questionnaire_catalog_exposes_no_numeric_scores(self):
        catalog = subjective_catalog()
        self.assertNotIn("score", str(catalog))

    def test_export_matches_bonus_and_capped_api_values(self):
        value = analysis([session(statuses=("yes",) * 30)])
        row = build_statistics(value, True)["rows"][0]
        wb = load_workbook(BytesIO(build_statistics_workbook(value, True)))
        sheet = wb["分数统计（试算）"]
        exported = dict(zip(next(sheet.values), list(sheet.values)[1]))
        for label, key in [("评审过程表现（30）", "objective_total"), ("专业价值贡献（70）", "subjective_total"),
                           ("超额意见奖励", "opinion_bonus"), ("对策奖励", "solution_bonus"), ("封顶前合计", "uncapped_total"), ("总分（100）", "total")]:
            self.assertEqual(row[key], exported[label])
        self.assertEqual(29, wb["阶段计分依据"]["M2"].value)

    def test_bonus_counts_opinions_with_or_without_solutions(self):
        row = self.row([session(statuses=("no", "yes", "no"))])
        self.assertEqual(2, row["opinion_bonus"])
        self.assertEqual(10, row["stages"]["TDR1"]["components"]["opinion"])

    def test_more_opinions_do_not_compare_against_other_experts(self):
        value = analysis([session(statuses=("yes",) * 5)])
        before = build_statistics(value, True)["rows"][0]
        value.experts.append(ExpertFacts("虚拟其他", [session(statuses=("yes",) * 100)], {}, {}))
        after = build_statistics(value, True)["rows"][0]
        self.assertEqual(before["opinion_bonus"], after["opinion_bonus"])
        self.assertEqual(before["stages"]["TDR1"]["components"]["opinion"], after["stages"]["TDR1"]["components"]["opinion"])
