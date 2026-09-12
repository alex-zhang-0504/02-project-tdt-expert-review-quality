from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from tdt_scoring.build_info import BUILD_ID
from tdt_scoring.service import ScoringService
from tdt_scoring.submission import (
    RULE_VERSION,
    SCHEMA_VERSION,
    build_dimension_one_workbook,
    load_dimension_one_workbook,
)
from tests.workbook_factory import build_v04_workbook


class DimensionOneSubmissionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.service = ScoringService()

    def _analysis(self, project: str, filename: str):
        return self.service.import_local_bytes(
            build_v04_workbook(
                [
                    {
                        "stage": "TDR3",
                        "signoffs": [
                            {
                                "role": "评审主席",
                                "reviewer": "虚拟评审人甲",
                                "conclusion": "Go",
                                "opinion": "针对链路预算，建议补充边缘场景仿真数据。",
                            },
                            {"role": "领域专家", "reviewer": "虚拟评审人乙", "conclusion": "-", "opinion": ""},
                            {"role": "测试专家", "reviewer": "虚拟评审人丙", "conclusion": "-", "opinion": ""},
                        ],
                    }
                ],
                project=project,
            ),
            filename,
        )

    def _submission(self, analysis, manager_id: str, manager_name: str = "虚拟项目经理") -> bytes:
        return build_dimension_one_workbook(
            analysis,
            package_kind="manager_submission",
            batch_id="2025年度评分",
            manager_id=manager_id,
            manager_name=manager_name,
            revision=1,
            product_version="v0.1",
            build_id=BUILD_ID,
        )

    def test_submission_contains_readable_sheets_and_verified_payload(self) -> None:
        analysis = self._analysis("虚拟大TDT-子任务甲-B260001", "B260001.xlsx")

        content = self._submission(analysis, "PM01")
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True)
        package = load_dimension_one_workbook(content, "PM01.xlsx")

        self.assertEqual(
            [
                "00_提交信息",
                "01_项目清单",
                "02_场次事实",
                "03_意见证据",
                "04_分阶段统计",
                "05_全部阶段汇总",
                "06_异常清单",
                "_manifest",
                "_payload",
            ],
            workbook.sheetnames,
        )
        self.assertEqual("veryHidden", workbook["_manifest"].sheet_state)
        self.assertEqual("veryHidden", workbook["_payload"].sheet_state)
        self.assertEqual(SCHEMA_VERSION, workbook.properties.subject)
        self.assertEqual("facts-v0.6", RULE_VERSION)
        self.assertEqual(RULE_VERSION, package.rule_version)
        self.assertEqual("试算", workbook["00_提交信息"]["B8"].value)
        headers = [cell.value for cell in workbook["04_分阶段统计"][1]]
        self.assertIn("TDR1 意见提出率", headers)
        self.assertIn("代理率", headers)
        self.assertEqual("PM01", package.manager_id)
        self.assertEqual(["B260001"], package.project_codes)

    def test_two_manager_submissions_recompute_one_annual_cohort(self) -> None:
        first = self._analysis("虚拟大TDT-子任务甲-B260001", "B260001.xlsx")
        second = self._analysis("虚拟大TDT-子任务乙-B260002", "B260002.xlsx")

        merged = self.service.merge_dimension_one_submissions(
            [
                (self._submission(first, "PM01"), "PM01.xlsx"),
                (self._submission(second, "PM02"), "PM02.xlsx"),
            ],
            expected_manager_count=2,
            expected_project_count=2,
        )

        self.assertFalse(any(issue.severity == "error" for issue in merged.issues))
        self.assertEqual(2, len({session.project_code for session in merged.sessions}))
        self.assertEqual(3, len(merged.experts))
        self.assertEqual(
            ["B260001-TDR3", "B260002-TDR3"],
            [s.review_id for s in merged.experts[0].sessions],
        )

    def test_sixteen_manager_submissions_recompute_one_annual_cohort(self) -> None:
        submissions = []
        for index in range(1, 17):
            project_code = f"B26{index:04d}"
            analysis = self._analysis(
                f"虚拟大TDT-子任务{index:02d}-{project_code}",
                f"{project_code}.xlsx",
            )
            submissions.append(
                (
                    self._submission(
                        analysis,
                        f"PM{index:02d}",
                        f"虚拟项目经理{index:02d}",
                    ),
                    f"PM{index:02d}.xlsx",
                )
            )

        merged = self.service.merge_dimension_one_submissions(
            submissions,
            expected_manager_count=16,
            expected_project_count=16,
        )

        self.assertFalse(any(issue.severity == "error" for issue in merged.issues))
        self.assertEqual(16, len({session.project_code for session in merged.sessions}))
        self.assertEqual(3, len(merged.experts))
        self.assertEqual(16, len([s.review_id for s in merged.experts[0].sessions]))

    def test_duplicate_project_across_managers_blocks_scores(self) -> None:
        analysis = self._analysis("虚拟大TDT-子任务甲-B260001", "B260001.xlsx")

        merged = self.service.merge_dimension_one_submissions(
            [
                (self._submission(analysis, "PM01"), "PM01.xlsx"),
                (self._submission(analysis, "PM02"), "PM02.xlsx"),
            ],
            expected_manager_count=2,
            expected_project_count=1,
        )

        self.assertEqual([], merged.experts)
        self.assertIn(
            "submission_project_owner_conflict",
            {issue.code for issue in merged.issues},
        )
        self.assertIn("submission_stage_duplicate", {issue.code for issue in merged.issues})

    def test_payload_edit_is_rejected(self) -> None:
        analysis = self._analysis("虚拟大TDT-子任务甲-B260001", "B260001.xlsx")
        content = self._submission(analysis, "PM01")
        workbook = load_workbook(BytesIO(content))
        workbook["_payload"]["B1"] = str(workbook["_payload"]["B1"].value) + "x"
        output = BytesIO()
        workbook.save(output)

        with self.assertRaisesRegex(ValueError, "校验失败"):
            load_dimension_one_workbook(output.getvalue(), "PM01-edited.xlsx")


if __name__ == "__main__":
    unittest.main()
